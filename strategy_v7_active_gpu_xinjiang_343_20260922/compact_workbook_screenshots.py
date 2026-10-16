#!/usr/bin/env python3
"""Convert floating evidence images to compact, cell-sized thumbnails."""

import argparse
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PilImage


MAX_WIDTH = 180
MAX_HEIGHT = 96


def thumbnail(image: ExcelImage) -> ExcelImage:
    with PilImage.open(BytesIO(image._data())) as source:
        source.thumbnail((MAX_WIDTH, MAX_HEIGHT), PilImage.Resampling.LANCZOS)
        if source.mode not in ("RGB", "RGBA"):
            source = source.convert("RGBA")
        output = BytesIO()
        source.save(output, format="PNG", optimize=True)
    return ExcelImage(BytesIO(output.getvalue()))


def row_from_anchor(image: ExcelImage) -> int:
    anchor = image.anchor
    if hasattr(anchor, "_from"):
        return anchor._from.row + 1
    if isinstance(anchor, str):
        digits = "".join(character for character in anchor if character.isdigit())
        return int(digits)
    raise ValueError("image has no usable row anchor")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    workbook = load_workbook(args.input, data_only=False)
    converted = 0
    for sheet in workbook.worksheets:
        if sheet.title == "汇总统计":
            continue
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        if "证明截图" not in headers:
            continue
        column = headers.index("证明截图") + 1
        original_images = list(sheet._images)
        sheet._images = []
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 27
        for image in original_images:
            row = row_from_anchor(image)
            compact = thumbnail(image)
            # Two-cell anchor covers exactly the screenshot cell, so Excel
            # moves and scales the picture with its row and column.
            compact.anchor = TwoCellAnchor(
                editAs="twoCell",
                _from=AnchorMarker(col=column - 1, row=row - 1, colOff=pixels_to_EMU(2), rowOff=pixels_to_EMU(2)),
                to=AnchorMarker(col=column, row=row, colOff=pixels_to_EMU(0), rowOff=pixels_to_EMU(0)),
            )
            sheet.add_image(compact)
            sheet.row_dimensions[row].height = 76
            converted += 1
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    workbook.close()
    print("converted=%d output=%s" % (converted, args.output))


if __name__ == "__main__":
    main()
