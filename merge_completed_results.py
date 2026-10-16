# -*- coding: utf-8 -*-
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


ROOT = Path("/home/lixiao/llm-detect/TEST/data/rescan_requested_20260813")
OUTPUT = ROOT / "已完成11家结果汇总.xlsx"


def copy_sheet(source, target, source_start_row, target_start_row):
    target.sheet_format.defaultColWidth = source.sheet_format.defaultColWidth
    target.sheet_format.defaultRowHeight = source.sheet_format.defaultRowHeight
    target.sheet_view.showGridLines = source.sheet_view.showGridLines

    for key, dimension in source.column_dimensions.items():
        if key not in target.column_dimensions:
            target.column_dimensions[key] = copy(dimension)
    for key, dimension in source.row_dimensions.items():
        if key >= source_start_row:
            target.row_dimensions[key - source_start_row + target_start_row] = copy(dimension)

    for row in source.iter_rows(min_row=source_start_row):
        for cell in row:
            new_cell = target.cell(cell.row - source_start_row + target_start_row, cell.column)
            new_cell.value = cell.value
            if cell.has_style:
                new_cell._style = copy(cell._style)
            if cell.number_format:
                new_cell.number_format = cell.number_format
            if cell.alignment:
                new_cell.alignment = copy(cell.alignment)
            if cell.protection:
                new_cell.protection = copy(cell.protection)
            if cell.hyperlink:
                new_cell._hyperlink = copy(cell.hyperlink)
            if cell.comment:
                new_cell.comment = copy(cell.comment)

    for merged_range in source.merged_cells.ranges:
        if merged_range.min_row < source_start_row:
            continue
        target.merge_cells(
            start_row=merged_range.min_row - source_start_row + target_start_row,
            start_column=merged_range.min_col,
            end_row=merged_range.max_row - source_start_row + target_start_row,
            end_column=merged_range.max_col,
        )
    return target_start_row + source.max_row - source_start_row + 1


def find_header(sheet):
    ip_column = None
    header_row = None
    for row in range(1, min(sheet.max_row, 10) + 1):
        for column in range(1, sheet.max_column + 1):
            value = str(sheet.cell(row, column).value or "").strip()
            if value == "IP地址":
                header_row, ip_column = row, column
                break
        if ip_column:
            break
    return header_row, ip_column


def count_records(sheet):
    header_row, ip_column = find_header(sheet)
    if not ip_column:
        return 0
    return sum(
        1
        for row in range(header_row + 1, sheet.max_row + 1)
        if sheet.cell(row, ip_column).value not in (None, "")
    )


def category_for(sheet_name):
    if "高" in sheet_name:
        return "高"
    if "中" in sheet_name:
        return "中"
    if "低" in sheet_name:
        return "低"
    return "其他"


def main():
    sources = []
    for directory in sorted(path for path in ROOT.iterdir() if path.is_dir()):
        result = directory / f"{directory.name}_llm_最终版.xlsx"
        if (directory / "COMPLETE").is_file() and result.is_file():
            sources.append((directory.name, result))

    if not sources:
        raise RuntimeError("未发现带 COMPLETE 标识的最终结果文件")

    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = workbook.create_sheet("汇总统计", 0)
    summary.append(["企业名称", "GPU可能性高", "GPU可能性中", "其他结果", "结果合计"])

    for enterprise, result in sources:
        source_workbook = load_workbook(result, data_only=False)
        counts = {"高": 0, "中": 0, "低": 0, "其他": 0}
        for source_sheet in source_workbook.worksheets:
            category = category_for(source_sheet.title)
            if category != "低":
                counts[category] += count_records(source_sheet)
        summary.append([
            enterprise,
            counts["高"],
            counts["中"],
            counts["其他"],
            counts["高"] + counts["中"] + counts["其他"],
        ])

        target_sheet = workbook.create_sheet(enterprise)
        next_row = 1
        is_first_table = True
        for source_sheet in source_workbook.worksheets:
            if category_for(source_sheet.title) == "低":
                continue
            header_row, _ = find_header(source_sheet)
            if header_row is None:
                continue
            source_start_row = header_row if is_first_table else header_row + 1
            next_row = copy_sheet(source_sheet, target_sheet, source_start_row, next_row)
            is_first_table = False
        if not is_first_table:
            target_sheet.auto_filter.ref = (
                f"A1:{get_column_letter(target_sheet.max_column)}{target_sheet.max_row}"
            )
            target_sheet.freeze_panes = "A2"
        source_workbook.close()

    summary.append([
        "合计",
        f"=SUM(B2:B{summary.max_row})",
        f"=SUM(C2:C{summary.max_row})",
        f"=SUM(D2:D{summary.max_row})",
        f"=SUM(E2:E{summary.max_row})",
    ])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    for column, width in {"A": 22, "B": 16, "C": 16, "D": 14, "E": 14}.items():
        summary.column_dimensions[column].width = width
    summary.freeze_panes = "A2"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(OUTPUT)
    workbook.close()
    print(f"已汇总 {len(sources)} 家企业：{OUTPUT}")
    print("\n".join(name for name, _ in sources))


if __name__ == "__main__":
    main()
