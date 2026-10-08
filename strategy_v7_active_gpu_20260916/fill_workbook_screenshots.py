#!/usr/bin/env python3
"""Fill blank screenshot cells in a workbook from an exported image manifest."""

import argparse
import json
import zipfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--images", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    workbook = load_workbook(args.input, data_only=False)
    inserted = 0
    with zipfile.ZipFile(args.images) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        for item in manifest:
            sheet = workbook[item["sheet"]]
            header = [str(cell.value or "").strip() for cell in sheet[1]]
            screenshot_column = header.index("证明截图") + 1
            target_row = int(item["row"])
            if any(
                hasattr(image.anchor, "_from")
                and image.anchor._from.row + 1 == target_row
                and image.anchor._from.col + 1 == screenshot_column
                for image in sheet._images
            ):
                continue
            image = ExcelImage(BytesIO(archive.read(item["image"])))
            image.anchor = sheet.cell(target_row, screenshot_column).coordinate
            sheet.add_image(image)
            if sheet.row_dimensions[target_row].height is None:
                sheet.row_dimensions[target_row].height = max(90, image.height * 0.75)
            inserted += 1
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    workbook.close()
    print("inserted=%d output=%s" % (inserted, args.output))


if __name__ == "__main__":
    main()
