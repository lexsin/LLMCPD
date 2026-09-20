#!/usr/bin/env python3

import json
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image

from screenshot_evidence import evidence_key, fill_workbook


with tempfile.TemporaryDirectory() as directory:
    root = Path(directory)
    source = root / "source.xlsx"
    output = root / "output.xlsx"
    asset = root / "evidence.png"
    manifest = root / "manifest.json"
    Image.new("RGB", (400, 180), "white").save(asset)
    book = Workbook()
    sheet = book.active
    sheet.append(["IP地址", "端口号", "协议类型", "访问链接", "存在GPU算力的说明", "证明截图"])
    sheet.append(["192.0.2.10", "11434", "http", "http://192.0.2.10:11434/api/ps", "size_vram=1024", ""])
    book.save(source)
    manifest.write_text(json.dumps({
        evidence_key("192.0.2.10", "11434", "http"): {
            "path": str(asset), "url": "http://192.0.2.10:11434/api/ps"
        }
    }), encoding="utf-8")
    records = fill_workbook(
        source, output, root / "assets", False, manifest_path=manifest
    )
    assert records[0][2] == "已填充预先截图"
    assert len(load_workbook(output).active._images) == 1

print("screenshot_precaptured_tests=passed")
