#!/usr/bin/env python3
"""Export source report images that can fill blank screenshot cells in a workbook."""

import argparse
import json
import re
import zipfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


def key_for(sheet, row, ip_col, port_col):
    ip = sheet.cell(row, ip_col).value
    if ip in (None, ""):
        return None
    port = sheet.cell(row, port_col).value if port_col else ""
    return str(ip), str(port)


def headers(sheet):
    values = [str(cell.value or "").strip() for cell in sheet[1]]
    return {
        "ip": values.index("IP地址") + 1 if "IP地址" in values else None,
        "port": values.index("端口号") + 1 if "端口号" in values else None,
        "screenshot": values.index("证明截图") + 1 if "证明截图" in values else None,
    }


def image_rows(sheet):
    return {
        image.anchor._from.row + 1: image
        for image in sheet._images
        if hasattr(image.anchor, "_from")
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--target", type=Path, required=True)
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    target = load_workbook(args.target, data_only=False)
    manifest = []
    with zipfile.ZipFile(args.output, "w", zipfile.ZIP_DEFLATED) as archive:
        for target_sheet in target.worksheets:
            if target_sheet.title == "汇总统计":
                continue
            info = headers(target_sheet)
            if not info["ip"] or not info["screenshot"]:
                continue
            source_file = args.source_root / target_sheet.title / (target_sheet.title + "_llm_最终版.xlsx")
            if not source_file.is_file():
                continue
            source = load_workbook(source_file, data_only=False)
            source_by_key = {}
            for source_sheet in source.worksheets:
                source_info = headers(source_sheet)
                if not source_info["ip"]:
                    continue
                for row, image in image_rows(source_sheet).items():
                    key = key_for(source_sheet, row, source_info["ip"], source_info["port"])
                    if key:
                        source_by_key[key] = image
            target_images = image_rows(target_sheet)
            for row in range(2, target_sheet.max_row + 1):
                if row in target_images:
                    continue
                key = key_for(target_sheet, row, info["ip"], info["port"])
                image = source_by_key.get(key)
                if not image:
                    continue
                image_name = "%s_%s.png" % (re.sub(r"[^\w.-]+", "_", target_sheet.title), row)
                archive.writestr(image_name, image._data())
                manifest.append({"sheet": target_sheet.title, "row": row, "image": image_name})
            source.close()
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    target.close()
    print("candidates=%d output=%s" % (len(manifest), args.output))


if __name__ == "__main__":
    main()
