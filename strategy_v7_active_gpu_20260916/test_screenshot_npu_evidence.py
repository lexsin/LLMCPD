import csv,json,tempfile,unittest,sys,io
from pathlib import Path
from unittest.mock import patch
from PIL import Image
import screenshot_evidence as s

METRIC='npu_chip_info_hbm_total_memory'
NOTE='/metrics: '+METRIC+'{id="0",model_name="910B4-1-Ascend-V1"}=65536'
RESPONSE=(200,'text/plain',(METRIC+'{id="0",model_name="910B4-1-Ascend-V1",vdie_id="ABC"} 65536\n').encode())
class NPUScreenshotTest(unittest.TestCase):
 def test_selection_and_derived_url(self):
  note=s.select_evidence_note('未见GPU显存',NOTE,'GPU算力可能性未知','NPU算力可能性高')
  self.assertEqual(note,NOTE)
  self.assertEqual(s.evidence_kind('NPU显存：'+note),'npu')
  self.assertEqual(s.preferred_urls('http://192.0.2.1:8800/',note)[0],'http://192.0.2.1:8800/metrics')
  self.assertEqual(s.select_evidence_note('size_vram=123',NOTE,'高','高'),'size_vram=123')
 def test_exact_metric_after_many_lines(self):
  body=('gpu_power 100\n'*1000+'npu_chip_info_hbm_used_memory 3\n'*100+
        RESPONSE[2].decode()).encode()
  lines,found=s.response_card_lines('text/plain',body,NOTE)
  self.assertTrue(found)
  self.assertTrue(any(x.startswith(METRIC+'{')for x in lines))
  self.assertLess(len(lines),10)
  self.assertFalse(s.response_matches_note((200,'text/html',b'<html>NPU-Exporter</html>'),NOTE))
  self.assertFalse(s.response_matches_note((200,'text/plain',b'npu_chip_info_hbm_used_memory 8'),NOTE))
 def test_csv_unknown_gpu_and_stale_cache(self):
  with tempfile.TemporaryDirectory()as td:
   d=Path(td);source=d/'source.csv';assets=d/'assets';assets.mkdir();mp=assets/'manifest.json'
   image=assets/'old.png';Image.new('RGB',(20,20),'white').save(image)
   key=s.evidence_key('192.0.2.1','8800','http')
   mp.write_text(json.dumps({key:{'path':str(image),'url':'http://192.0.2.1:8800/'}}))
   with source.open('w',newline='')as f:
    w=csv.DictWriter(f,fieldnames=['ip','port','protocol','gpu_likelihood','npu_likelihood','gpu_evidence','npu_evidence','link']);w.writeheader()
    w.writerow(dict(ip='192.0.2.1',port='8800',protocol='http',gpu_likelihood='未知',npu_likelihood='高',gpu_evidence='GPU未知',npu_evidence=NOTE,link='http://192.0.2.1:8800/'))
   calls=[]
   def fetch(url):calls.append(url);return RESPONSE
   with patch.object(s,'fetch_response',side_effect=fetch):
    count,records=s.capture_csv_evidence(source,assets,mp)
   self.assertEqual(count,1);self.assertEqual(calls,['http://192.0.2.1:8800/metrics'])
   self.assertTrue(s.reusable_precapture(records[key],NOTE,calls))
   with patch.object(s,'fetch_response',side_effect=AssertionError('must reuse matched cache')):
    count,_=s.capture_csv_evidence(source,assets,mp)
   self.assertEqual(count,0)
 def test_csv_no_matching_response_is_not_fake_evidence(self):
  with tempfile.TemporaryDirectory()as td:
   d=Path(td);source=d/'source.csv'
   with source.open('w',newline='')as f:
    w=csv.DictWriter(f,fieldnames=['ip','port','protocol','gpu_likelihood','npu_likelihood','npu_evidence','link']);w.writeheader()
    w.writerow(dict(ip='192.0.2.1',port='8800',protocol='http',gpu_likelihood='未知',npu_likelihood='高',npu_evidence=NOTE,link='http://192.0.2.1:8800/'))
   with patch.object(s,'fetch_response',return_value=(200,'text/html',b'<html>MinIO Console</html>')):
    count,records=s.capture_csv_evidence(source,d/'assets',d/'manifest.json')
   self.assertEqual(count,0);self.assertFalse(records)
 def test_workbook_uses_npu_note_not_gpu_unknown(self):
  from openpyxl import Workbook,load_workbook
  with tempfile.TemporaryDirectory()as td:
   d=Path(td);src=d/'src.xlsx';out=d/'out.xlsx';wb=Workbook();ws=wb.active
   ws.append(['IP地址','端口号','协议类型','访问链接','存在GPU算力的说明','存在GPU算力概率','NPU算力说明','NPU算力可能性','证明截图'])
   ws.append(['192.0.2.1',8800,'http','http://192.0.2.1:8800/','未见GPU显存','GPU算力可能性未知',NOTE,'NPU算力可能性高',''])
   wb.save(src);calls=[]
   def fetch(url):calls.append(url);return RESPONSE
   with patch.object(s,'fetch_response',side_effect=fetch):
    records=s.fill_workbook(src,out,d/'assets',False)
   self.assertEqual(calls,['http://192.0.2.1:8800/metrics'])
   self.assertEqual(len(load_workbook(out).active._images),1)
   self.assertEqual(records[0][3],calls[0])
 def test_stream_retains_npu_only_tail_without_newline(self):
  import io
  class Response(io.BytesIO):
   status=200
   headers={'Content-Type':'text/plain'}
  tail=b'npu_device_count{device="Ascend910"} 8'
  with patch.object(s,'urlopen',return_value=Response(b'ordinary_counter 1\n'*10000+tail)):
   result=s.fetch_response('http://192.0.2.1:8800/metrics')
  self.assertIn(tail,result[2])
 def test_unknown_npu_message_does_not_replace_gpu_note(self):
  self.assertEqual(s.select_evidence_note('', '未获得有效NPU证据','低','未知'),'')

if __name__=='__main__':unittest.main()
