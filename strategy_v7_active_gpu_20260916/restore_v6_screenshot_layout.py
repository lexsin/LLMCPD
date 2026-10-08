#!/usr/bin/env python3
"""Apply the v5/v6 evidence-image layout without recompressing screenshots."""

import argparse
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor


DISPLAY_WIDTH = 400
DISPLAY_HEIGHT = 180
ROW_HEIGHT = 140
COLUMN_WIDTH = 62


def row_from_anchor(image):
    anchor = image.anchor
    if hasattr(anchor, "_from"):
        return anchor._from.row + 1
    if isinstance(anchor, str):
        return int("".join(character for character in anchor if character.isdigit()))
    raise ValueError("image has no usable row anchor")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    workbook = load_workbook(args.input, data_only=False)
    restored = 0
    for sheet in workbook.worksheets:
        if sheet.title == "汇总统计":
            continue
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        if "证明截图" not in headers:
            continue
        column = headers.index("证明截图") + 1
        original_images = list(sheet._images)
        sheet._images = []
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = COLUMN_WIDTH
        for image in original_images:
            row = row_from_anchor(image)
            restored_image = ExcelImage(BytesIO(image._data()))
            restored_image.width = DISPLAY_WIDTH
            restored_image.height = DISPLAY_HEIGHT
            restored_image.anchor = TwoCellAnchor(
                editAs="oneCell",
                _from=AnchorMarker(col=column - 1, colOff=0, row=row - 1, rowOff=0),
                to=AnchorMarker(col=column, colOff=0, row=row, rowOff=0),
            )
            sheet.add_image(restored_image)
            sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 0, ROW_HEIGHT)
            restored += 1
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    workbook.close()
    print("restored=%d output=%s" % (restored, args.output))


if __name__ == "__main__":
    main()
