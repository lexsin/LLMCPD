# -*- coding: utf-8 -*-
"""Merge completed enterprise reports, excluding rows with low GPU likelihood."""
import argparse
import re
from copy import copy, deepcopy
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reported_ip_lookup import add_reported_column, load_reported_intervals


ROOT = Path("/home/lixiao/llm-detect/TEST/data/rescan_all_isolated_v4_20260907")
OUTPUT = ROOT / "已完成企业结果汇总.xlsx"
SUMMARY_HEADERS = ["企业名称", "GPU可能性高", "GPU可能性中", "其他结果", "结果合计"]
GPU_HEADERS = {"存在GPU算力概率", "存在GPU算力可能性", "GPU算力可能性", "GPU可能性"}


def find_header(sheet):
    for row in range(1, min(sheet.max_row, 10) + 1):
        for column in range(1, sheet.max_column + 1):
            if str(sheet.cell(row, column).value or "").strip() == "IP地址":
                return row, column
    return None, None


def category_for(sheet_name):
    for category in ("高", "中", "低"):
        if category in sheet_name:
            return category
    return "其他"


def gpu_category(value):
    value = re.sub(r"\s+", "", str(value or "")).lower()
    aliases = {
        "高": {"高", "high", "gpu算力可能性高", "gpu可能性高"},
        "中": {"中", "中等", "medium", "gpu算力可能性中", "gpu算力可能性中等", "gpu可能性中"},
        "低": {"低", "较低", "low", "gpu算力可能性低", "gpu算力可能性较低", "gpu可能性低"},
    }
    for category, values in aliases.items():
        if value in values:
            return category
    return "其他"


def row_categories(sheet):
    header_row, ip_column = find_header(sheet)
    if header_row is None:
        return {}
    gpu_column = next(
        (cell.column for cell in sheet[header_row] if str(cell.value or "").strip() in GPU_HEADERS),
        None,
    )
    return {
        row: (
            gpu_category(sheet.cell(row, gpu_column).value)
            if gpu_column else category_for(sheet.title)
        )
        for row in range(header_row + 1, sheet.max_row + 1)
        if sheet.cell(row, ip_column).value not in (None, "")
    }


def count_categories(sheet, exclude_low=True):
    counts = {"高": 0, "中": 0, "低": 0, "其他": 0}
    for category in row_categories(sheet).values():
        if not (exclude_low and category == "低"):
            counts[category] += 1
    return counts


def count_records(sheet):
    return sum(count_categories(sheet).values())


def copy_sheet(source, target, source_start_row, target_start_row):
    """Compact selected rows and move each retained screenshot with its row."""
    categories = row_categories(source)
    kept_rows = [
        row for row in range(source_start_row, source.max_row + 1)
        if categories.get(row) != "低"
    ]
    row_map = {row: target_start_row + i for i, row in enumerate(kept_rows)}
    target.sheet_format.defaultColWidth = source.sheet_format.defaultColWidth
    target.sheet_format.defaultRowHeight = source.sheet_format.defaultRowHeight
    target.sheet_view.showGridLines = source.sheet_view.showGridLines

    for key, dimension in source.column_dimensions.items():
        if key not in target.column_dimensions:
            target.column_dimensions[key] = copy(dimension)
    for row in kept_rows:
        if row in source.row_dimensions:
            dimension = copy(source.row_dimensions[row])
            dimension.index = row_map[row]
            target.row_dimensions[row_map[row]] = dimension
        for cell in source[row]:
            new_cell = target.cell(row_map[row], cell.column, cell.value)
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.fill = copy(cell.fill)
                new_cell.border = copy(cell.border)
            if cell.number_format:
                new_cell.number_format = cell.number_format
            if cell.alignment:
                new_cell.alignment = copy(cell.alignment)
            if cell.protection:
                new_cell.protection = copy(cell.protection)
            if cell.hyperlink:
                hyperlink = copy(cell.hyperlink)
                hyperlink.ref = new_cell.coordinate
                new_cell._hyperlink = hyperlink
            if cell.comment:
                new_cell.comment = copy(cell.comment)

    for merged_range in source.merged_cells.ranges:
        if merged_range.min_row < source_start_row:
            continue
        rows = list(range(merged_range.min_row, merged_range.max_row + 1))
        retained = [row for row in rows if row in row_map]
        if not retained:
            continue
        if len(retained) != len(rows):
            raise ValueError(f"过滤行穿过合并单元格，无法安全保留：{source.title}!{merged_range}")
        target.merge_cells(
            start_row=row_map[merged_range.min_row],
            start_column=merged_range.min_col,
            end_row=row_map[merged_range.max_row],
            end_column=merged_range.max_col,
        )

    for image in source._images:
        anchor = image.anchor
        if not hasattr(anchor, "_from"):
            raise ValueError(f"无法定位截图所属行：{source.title}")
        source_row = anchor._from.row + 1
        if source_row not in row_map:
            continue
        copied_image = ExcelImage(BytesIO(image._data()))
        copied_image.width = image.width
        copied_image.height = image.height
        copied_anchor = deepcopy(anchor)
        row_delta = row_map[source_row] - source_row
        copied_anchor._from.row += row_delta
        if hasattr(copied_anchor, "to"):
            copied_anchor.to.row += row_delta
        copied_image.anchor = copied_anchor
        target.add_image(copied_image)
    return target_start_row + len(kept_rows)


def new_workbook():
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = workbook.create_sheet("汇总统计", 0)
    summary.append(SUMMARY_HEADERS)
    return workbook, summary


def append_summary(summary, enterprise, counts):
    summary.append([
        enterprise, counts["高"], counts["中"], counts["其他"],
        counts["高"] + counts["中"] + counts["其他"],
    ])


def finish_workbook(workbook, summary, output):
    last_data_row = summary.max_row
    summary.append(["合计"] + [
        f"=SUM({column}2:{column}{last_data_row})" for column in ("B", "C", "D", "E")
    ])
    for row in (summary[1], summary[summary.max_row]):
        for cell in row:
            cell.font = Font(bold=True)
    for column, width in {"A": 22, "B": 16, "C": 16, "D": 14, "E": 14}.items():
        summary.column_dimensions[column].width = width
    summary.freeze_panes = "A2"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output = Path(output)
    temporary = output.with_name(output.stem + ".tmp.xlsx")
    try:
        workbook.save(temporary)
        temporary.replace(output)
    finally:
        workbook.close()


def configure_detail(sheet):
    if find_header(sheet)[0] is not None:
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        sheet.freeze_panes = "A2"



def main(root=ROOT, output=OUTPUT, database=None, reported_table="第二批上报算力"):
    contains = None
    interval_count = 0
    if database:
        contains, interval_count = load_reported_intervals(database, reported_table)
    sources = []
    for directory in sorted(path for path in root.iterdir() if path.is_dir()):
        result = directory / f"{directory.name}_llm_最终版.xlsx"
        if (directory / "COMPLETE").is_file() and result.is_file():
            sources.append((directory.name, result))
    if not sources:
        raise RuntimeError("未发现带 COMPLETE 标识的最终结果文件")

    workbook, summary = new_workbook()
    reported_yes = reported_no = excluded = 0
    for enterprise, result in sources:
        source_workbook = load_workbook(result, data_only=False)
        counts = {"高": 0, "中": 0, "低": 0, "其他": 0}
        for source_sheet in source_workbook.worksheets:
            for category, value in count_categories(source_sheet).items():
                counts[category] += value
            excluded += count_categories(source_sheet, exclude_low=False)["低"]
        append_summary(summary, enterprise, counts)

        target_sheet = workbook.create_sheet(enterprise)
        next_row = 1
        is_first_table = True
        for source_sheet in source_workbook.worksheets:
            header_row, _ = find_header(source_sheet)
            if header_row is None:
                continue
            source_start_row = header_row if is_first_table else header_row + 1
            next_row = copy_sheet(source_sheet, target_sheet, source_start_row, next_row)
            is_first_table = False
        if not is_first_table:
            if contains:
                yes, no = add_reported_column(target_sheet, contains)
                reported_yes += yes
                reported_no += no
            configure_detail(target_sheet)
        source_workbook.close()
    finish_workbook(workbook, summary, output)
    print(f"已汇总 {len(sources)} 家企业，排除GPU低可能性 {excluded} 条：{output}")
    if contains:
        print(
            f"是否上报: 是={reported_yes} 否={reported_no} "
            f"合并IP区间={interval_count} 表={reported_table}"
        )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="汇总已完成企业的LLM结果，排除GPU低可能性")
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--database", type=Path)
    parser.add_argument("--reported-table", default="第二批上报算力")
    args = parser.parse_args()
    main(
        args.root,
        args.output or args.root / "已完成企业结果汇总.xlsx",
        args.database,
        args.reported_table,
    )
