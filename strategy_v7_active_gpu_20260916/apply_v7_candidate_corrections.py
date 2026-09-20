#!/usr/bin/env python3
"""Apply the explicitly approved v7 candidate-logic corrections to a workbook."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

from filter_valid_gpu_screenshots import headers, refresh_summary
from screenshot_evidence import (
    add_metadata,
    attach_image,
    capture_response_card,
    detect_font,
    fetch_response,
    remove_image_at,
)


TARGET_STATUSES = {"should_low_filtered", "needs_review_high"}


def retained_models(profiles):
    names = []
    for profile in profiles:
        if profile.get("origin") != "local":
            continue
        kind = profile.get("kind")
        size = profile.get("parameter_b")
        if kind in {"embedding", "reranker"}:
            continue
        if kind == "generative" and isinstance(size, (int, float)) and size < 10:
            continue
        name = str(profile.get("name") or "").strip()
        if name and name not in names:
            names.append(name)
    return names[:4]


def correction_note(item):
    if item["status"] == "should_low_filtered":
        names = "、".join(profile["name"] for profile in item["models"][:4]) or "返回模型"
        return (
            "/v1/models响应中仅出现%s（embedding或reranker模型）；"
            "该类模型已从本地GPU候选中排除，且当前响应未出现GPU设备、显存占用或运行指标，"
            "因此GPU算力可能性低。" % names
        ), "GPU算力可能性低"
    models = retained_models(item["models"])
    if not models:
        raise ValueError("medium correction has no retained local model")
    return (
        "/v1/models响应中包含可保留本地模型标识%s；"
        "当前响应未出现GPU设备、显存占用或运行指标，因此GPU算力可能性中等。" % "、".join(models)
    ), "GPU算力可能性中等"


def main():
    parser = argparse.ArgumentParser(description="根据v7候选规则调整等级、说明与截图")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--audit", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--assets", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()

    audit = json.loads(args.audit.read_text(encoding="utf-8"))
    targets = [item for item in audit["rows"] if item["status"] in TARGET_STATUSES]
    if len(targets) != 6:
        raise RuntimeError("expected 6 approved corrections, got %d" % len(targets))

    # Fetch every response before touching the workbook so no partial output
    # can be produced if a target has changed or become inaccessible.
    prepared = []
    for item in targets:
        response = fetch_response(item["url"])
        if response is None:
            raise RuntimeError("cannot refresh current evidence for %s:%s" % (item["ip"], item["port"]))
        note, probability = correction_note(item)
        prepared.append((item, response, note, probability))

    workbook = load_workbook(args.input)
    args.assets.mkdir(parents=True, exist_ok=True)
    font = str(detect_font())
    changes = []
    for item, response, note, probability in prepared:
        sheet = workbook[item["sheet"]]
        columns = headers(sheet)
        row = item["row"]
        if str(sheet.cell(row, columns["IP地址"]).value or "") != item["ip"]:
            raise RuntimeError("row changed for %s:%s" % (item["ip"], item["port"]))
        if str(sheet.cell(row, columns["端口号"]).value or "") != item["port"]:
            raise RuntimeError("port changed for %s:%s" % (item["ip"], item["port"]))
        image_path = args.assets / ("%s_%s_%s.png" % (sheet.title, row, item["status"]))
        capture_response_card(item["url"], response, image_path, font, note)
        add_metadata(image_path, item["url"], font)
        remove_image_at(sheet, row, columns["证明截图"])
        attach_image(sheet, image_path, row, columns["证明截图"])
        sheet.cell(row, columns["存在GPU算力概率"]).value = probability
        sheet.cell(row, columns["存在GPU算力的说明"]).value = note
        changes.append({
            "sheet": sheet.title, "row": row, "ip": item["ip"], "port": item["port"],
            "status": item["status"], "new_probability": probability, "new_note": note, "url": item["url"],
        })

    refresh_summary(workbook)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(args.output)
    args.report.write_text(json.dumps(changes, ensure_ascii=False, indent=2), encoding="utf-8")
    print("corrected=%d output=%s" % (len(changes), args.output))


if __name__ == "__main__":
    main()
