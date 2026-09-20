import asyncio
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook
import scan_llm as scan
import npu_discovery as npu
import finalize_llm_results as report


def classify(body, path="/metrics", status=200, tool="", llm="否"):
    state = scan.TargetState(ip="192.0.2.1", port="8000", protocol="http",
                             deploy_tool=tool, is_llm=llm)
    state.probes[path] = scan.ProbeResult(status=status, body=body)
    scan._apply_service_classification(state, scan.get_default_config())
    return state


SAMPLE = 'npu_chip_info_utilization{id="0",model_name="910A-Ascend-V1",vdie_id="chip-0"} 0'


class NpuTests(unittest.TestCase):
    def test_idle_device(self):
        state = classify(SAMPLE)
        self.assertEqual(state.npu_likelihood, "高")
        self.assertEqual(state.accelerator_type, "NPU")
        self.assertIn("chip-0", state.npu_evidence)
        self.assertEqual(state.is_llm, "否")

    def test_bad_samples(self):
        for body in [
            "# HELP npu_chip_info_utilization Ascend utilization\n# TYPE npu_chip_info_utilization gauge",
            'npu_chip_info_utilization{} 10',
            'npu_chip_info_utilization{id="0"} NaN',
            'npu_chip_info_utilization{id="0"} +Inf',
            'npu_chip_info_utilization{id="0"} -1',
            'npu_chip_info_utilization{id="0"} 101',
            'npu_chip_info_utilization{id="0" broken} 1',
            'npu_chip_info_hbm_total_memory{id="0"} 0',
            'npu_unknown_metric{id="0"} 1',
            '{"error": "Ascend not available"}',
        ]:
            with self.subTest(body=body):
                self.assertNotEqual(classify(body).npu_likelihood, "高")
        for status in (0, 404, 500):
            self.assertEqual(classify(SAMPLE, status=status).npu_likelihood, "未知")

    def test_foreign_instance(self):
        for host in ("192.0.2.2:9400", "remote.example:9400"):
            body = SAMPLE.replace('id="0"', 'id="0",instance="' + host + '"')
            state = classify(body)
            self.assertEqual(state.npu_likelihood, "未知")
            self.assertIn("instance", state.npu_evidence)
        body = SAMPLE.replace('id="0"', 'id="0",instance="192.0.2.1:9400"')
        self.assertEqual(classify(body).npu_likelihood, "高")

    def test_backend_not_gpu(self):
        state = classify('{"backend":"vllm-ascend"}', "/runtime", tool="vllm", llm="确认")
        self.assertEqual(state.npu_likelihood, "中")
        self.assertEqual(state.gpu_likelihood, "未知")
        self.assertEqual(state.accelerator_type, "NPU")
        generic = classify('{"backend":"vllm"}', "/runtime", tool="vllm", llm="确认")
        self.assertEqual(generic.npu_likelihood, "未知")
        self.assertEqual(generic.gpu_likelihood, "中")
        self.assertEqual(classify("<html>MindIE documentation</html>", "/").npu_likelihood, "低")

    def test_gpu_coexistence(self):
        state = classify('{"backend":"mindie"}', "/runtime", tool="vllm", llm="确认")
        state.probes["/metrics"] = scan.ProbeResult(
            status=200, body=SAMPLE + '\nDCGM_FI_DEV_GPU_UTIL{gpu="0",UUID="GPU-test"} 0')
        scan._apply_service_classification(state, scan.get_default_config())
        self.assertEqual(state.gpu_likelihood, "高")
        self.assertEqual(state.npu_likelihood, "高")
        self.assertEqual(state.accelerator_type, "GPU+NPU")

    def test_ascend_assignment_and_runtime_metrics(self):
        state = classify('{"accelerators":["ascend:0"]}', "/v1/models", tool="vllm", llm="确认")
        self.assertEqual(state.npu_likelihood, "中")
        self.assertEqual(state.gpu_likelihood, "未知")
        state.probes["/metrics"] = scan.ProbeResult(status=200, body="vllm:gpu_cache_usage_perc 0.4")
        scan._apply_service_classification(state, scan.get_default_config())
        self.assertEqual(state.gpu_likelihood, "未知")
        generic = classify('{"accelerators":["0"]}', "/v1/models")
        self.assertEqual(generic.npu_likelihood, "未知")

    def test_inventory_and_documentation(self):
        body = '{"devices":[{"device_name":"Ascend 910B","id":0}]}'
        self.assertEqual(classify(body, "/api/npu/devices").npu_likelihood, "高")
        for path in ("/openapi.json", "/v1/models", "/"):
            self.assertNotEqual(classify(body, path).npu_likelihood, "高")
        foreign = '{"instance":"192.0.2.2:8000","devices":[{"device_name":"Ascend 910B","id":0}]}'
        self.assertEqual(classify(foreign, "/api/npu/devices").npu_likelihood, "未知")

    def test_same_ip_annotation(self):
        source = classify(SAMPLE).to_row()
        target = scan.TargetState(ip="192.0.2.1", port="8001", is_llm="确认").to_row()
        rows = [source, target]
        self.assertGreater(npu.annotate_same_ip(rows), 0)
        self.assertEqual(target["npu_likelihood"], "未知")
        self.assertIn("same_ip_npu_port=8000", target["npu_evidence"])
        self.assertEqual(npu.annotate_same_ip(rows), 0)

    def test_checkpoint_and_legacy_csv(self):
        row = classify(SAMPLE).to_row()
        self.assertEqual(scan._row_to_state(row).to_row()["npu_evidence"], row["npu_evidence"])
        with tempfile.TemporaryDirectory() as tmp:
            csv = Path(tmp) / "old.csv"
            scan.write_csv_atomic(csv, ["ip", "port"], [{"ip": "192.0.2.2", "port": "80"}])
            scan.write_csv_rows(csv, [row], append=True)
            fields, rows = scan.read_csv_rows_preserve(csv)
            self.assertEqual(len(rows), 2)
            self.assertIn("npu_likelihood", fields)
            self.assertEqual(rows[1]["npu_likelihood"], "高")
            cp = Path(tmp) / "checkpoint.jsonl"
            scan.append_gpu_rescan_updates(cp, [row])
            self.assertEqual(scan.load_gpu_rescan_updates(cp)["192.0.2.1:8000"]["npu_likelihood"], "高")

    def test_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "npu.xlsx"
            stats = report.write_result(pd.DataFrame([classify(SAMPLE).to_row()]), out)
            self.assertEqual(stats["npu_high"], 1)
            book = load_workbook(out)
            self.assertIn("NPU算力资源", book.sheetnames)
            sheet = book["NPU算力资源"]
            self.assertEqual(sheet.max_row, 2)
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(headers[-3:], ["NPU算力可能性", "NPU算力说明", "设备类型"])
            self.assertIn("存在GPU算力概率", headers)
            book.close()
            old = report.transform(pd.DataFrame([{"ip":"192.0.2.5", "is_llm":"确认"}]))
            self.assertEqual(old.iloc[0]["NPU算力可能性"], "NPU算力可能性未知")
            report.write_result(pd.DataFrame(), Path(tmp) / "empty.xlsx")

    def test_read_only_paths(self):
        spec = {"paths": {p: {"get": {}} for p in (
            "/api/npu/devices", "/npu/status", "/npu/reset", "/npu/restart",
            "/npu/{id}", "/npu/allocate")}}
        self.assertEqual(set(scan._openapi_hardware_paths(json.dumps(spec))),
                         {"/npu/status", "/api/npu/devices"})

    def test_old_rescan_checkpoint_rechecked(self):
        async def probe(states, concurrency, cfg):
            for state in states:
                state.probes["/metrics"] = scan.ProbeResult(status=200, body=SAMPLE)
            return states
        async def no_probe(states, concurrency, cfg):
            return states
        with tempfile.TemporaryDirectory() as tmp:
            src, out, cp = (Path(tmp) / p for p in ("in.csv", "out.csv", "cp.jsonl"))
            old = {"ip":"192.0.2.1", "port":"8000", "protocol":"http", "is_llm":"否"}
            scan.write_csv_atomic(src, list(old), [old])
            scan.append_gpu_rescan_updates(cp, [old])
            with patch.object(scan, "phase_gpu_discovery", probe), patch.object(scan, "phase_gpu_openapi", no_probe):
                asyncio.run(scan.run_gpu_rescan(src, out, cp, scan.get_default_config(), 10, True))
            self.assertEqual(scan.read_csv_rows_preserve(out)[1][0]["npu_likelihood"], "高")


if __name__ == "__main__":
    unittest.main()
