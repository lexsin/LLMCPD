#!/usr/bin/env python3
"""Keep only high/medium GPU rows that have a current evidence screenshot."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


GPU_HEADERS = {"IP地址", "证明截图", "存在GPU算力概率"}


def headers(sheet):
    return {
        str(cell.value or "").strip(): column
        for column, cell in enumerate(sheet[1], start=1)
    }


def likelihood(value):
    text = str(value or "")
    if "高" in text:
        return "高"
    if "中" in text:
        return "中"
    return ""


def image_rows(sheet, proof_column):
    return {
        image.anchor._from.row + 1
        for image in sheet._images
        if hasattr(image.anchor, "_from")
        and image.anchor._from.col + 1 == proof_column
    }


def shift_images_after_row_delete(sheet, deleted_row):
    """Keep picture anchors bound to their record after ``delete_rows``."""
    deleted_index = deleted_row - 1
    for image in sheet._images:
        anchor = image.anchor
        if not hasattr(anchor, "_from"):
            continue
        if anchor._from.row > deleted_index:
            anchor._from.row -= 1
            anchor.to.row -= 1


def count_likelihood(sheet, probability_column):
    counts = {"高": 0, "中": 0}
    for row in range(2, sheet.max_row + 1):
        level = likelihood(sheet.cell(row, probability_column).value)
        if level:
            counts[level] += 1
    return counts


def refresh_summary(workbook):
    if "汇总统计" not in workbook.sheetnames:
        return []
    summary = workbook["汇总统计"]
    summary_headers = headers(summary)
    required = {"企业名称", "GPU可能性高", "GPU可能性中", "NPU可能性高", "结果合计"}
    if not required.issubset(summary_headers):
        raise RuntimeError("summary headers changed; refusing to update statistics")
    updated = []
    total_row = None
    for row in range(2, summary.max_row + 1):
        company = str(summary.cell(row, summary_headers["企业名称"]).value or "").strip()
        if company == "合计":
            total_row = row
            continue
        if company not in workbook.sheetnames:
            continue
        sheet = workbook[company]
        columns = headers(sheet)
        if not GPU_HEADERS.issubset(columns):
            continue
        counts = count_likelihood(sheet, columns["存在GPU算力概率"])
        npu_cell = summary.cell(row, summary_headers["NPU可能性高"])
        npu_high = npu_cell.value if isinstance(npu_cell.value, (int, float)) else 0
        summary.cell(row, summary_headers["GPU可能性高"]).value = counts["高"]
        summary.cell(row, summary_headers["GPU可能性中"]).value = counts["中"]
        summary.cell(row, summary_headers["结果合计"]).value = counts["高"] + counts["中"] + npu_high
        updated.append({"company": company, "high": counts["高"], "medium": counts["中"], "npu_high": npu_high})
    if total_row:
        high_total = sum(item["high"] for item in updated)
        medium_total = sum(item["medium"] for item in updated)
        # Some companies contain only NPU records and therefore do not have a
        # GPU likelihood column.  Preserve and sum their existing NPU counts.
        npu_total = sum(
            value if isinstance(value, (int, float)) else 0
            for row in range(2, summary.max_row + 1)
            if row != total_row
            for value in [summary.cell(row, summary_headers["NPU可能性高"]).value]
        )
        summary.cell(total_row, summary_headers["GPU可能性高"]).value = high_total
        summary.cell(total_row, summary_headers["GPU可能性中"]).value = medium_total
        summary.cell(total_row, summary_headers["NPU可能性高"]).value = npu_total
        summary.cell(total_row, summary_headers["结果合计"]).value = high_total + medium_total + npu_total
    return updated


def main():
    parser = argparse.ArgumentParser(description="删除无有效截图的GPU高/中记录")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()

    workbook = load_workbook(args.input)
    removed = []
    for sheet in workbook.worksheets:
        columns = headers(sheet)
        if not GPU_HEADERS.issubset(columns):
            continue
        proof_rows = image_rows(sheet, columns["证明截图"])
        rows_to_delete = [
            row for row in range(2, sheet.max_row + 1)
            if likelihood(sheet.cell(row, columns["存在GPU算力概率"]).value)
            and row not in proof_rows
        ]
        for row in reversed(rows_to_delete):
            removed.append({
                "sheet": sheet.title,
                "row_before_delete": row,
                "ip": str(sheet.cell(row, columns["IP地址"]).value or ""),
                "likelihood": likelihood(sheet.cell(row, columns["存在GPU算力概率"]).value),
            })
            sheet.delete_rows(row, 1)
            shift_images_after_row_delete(sheet, row)
        # A retained proof image must have enough row height after records are
        # compacted.  This is intentionally local to picture-bearing rows.
        for image in sheet._images:
            if hasattr(image.anchor, "_from"):
                image_row = image.anchor._from.row + 1
                sheet.row_dimensions[image_row].height = max(
                    sheet.row_dimensions[image_row].height or 0, 240
                )

    summary = refresh_summary(workbook)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(args.output)
    report = {
        "removed": removed,
        "removed_by_likelihood": {
            "高": sum(item["likelihood"] == "高" for item in removed),
            "中": sum(item["likelihood"] == "中" for item in removed),
        },
        "summary": summary,
    }
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print("removed=%d high=%d medium=%d" % (
        len(removed), report["removed_by_likelihood"]["高"], report["removed_by_likelihood"]["中"],
    ))


if __name__ == "__main__":
    main()
