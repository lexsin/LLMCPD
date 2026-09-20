#!/usr/bin/env python3
"""Passively audit retained GPU rows with v7's cloud/model-candidate rules."""
from __future__ import annotations

import argparse
import json
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from openpyxl import load_workbook

from refresh_remaining_gpu_evidence import direct_gpu_lines, risk_level
from scan_llm import (
    _is_explicit_cloud_model,
    _is_explicit_cloud_owner,
    _model_kind,
    _parse_parameter_billions,
)
from screenshot_evidence import fetch_response, preferred_urls, response_text


def headers(sheet):
    return {
        str(cell.value or "").strip(): column
        for column, cell in enumerate(sheet[1], start=1)
    }


def image_rows(sheet, image_column):
    return {
        image.anchor._from.row + 1
        for image in sheet._images
        if hasattr(image.anchor, "_from")
        and image.anchor._from.col + 1 == image_column
    }


def collect_targets(workbook):
    targets = []
    required = {"IP地址", "端口号", "访问链接", "证明截图", "存在GPU算力概率", "存在GPU算力的说明"}
    for sheet in workbook.worksheets:
        columns = headers(sheet)
        if not required.issubset(columns):
            continue
        pictures = image_rows(sheet, columns["证明截图"])
        for row in range(2, sheet.max_row + 1):
            level = risk_level(sheet.cell(row, columns["存在GPU算力概率"]).value)
            if not level or row not in pictures:
                continue
            targets.append({
                "sheet": sheet.title,
                "row": row,
                "ip": str(sheet.cell(row, columns["IP地址"]).value or ""),
                "port": str(sheet.cell(row, columns["端口号"]).value or ""),
                "level": level,
                "links": str(sheet.cell(row, columns["访问链接"]).value or ""),
                "note": str(sheet.cell(row, columns["存在GPU算力的说明"]).value or ""),
                "tool": str(sheet.cell(row, columns.get("模型部署工具", 0)).value or "") if columns.get("模型部署工具") else "",
            })
    return targets


def parse_models(text):
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        return []
    records = []
    if isinstance(data, dict):
        for key in ("data", "models"):
            items = data.get(key)
            if isinstance(items, list):
                records.extend(item for item in items if isinstance(item, dict))
    profiles = []
    seen = set()
    for item in records:
        name = str(item.get("id") or item.get("model_name") or item.get("name") or item.get("model") or "").strip()
        if not name or name.lower() in seen:
            continue
        seen.add(name.lower())
        details = item.get("details") if isinstance(item.get("details"), dict) else {}
        cloud = _is_explicit_cloud_model(name, item) or _is_explicit_cloud_owner(item.get("owned_by"))
        parameter = item.get("parameter_size") or item.get("parameters") or details.get("parameter_size")
        profiles.append({
            "name": name,
            "origin": "cloud" if cloud else "local",
            "kind": _model_kind(name),
            "parameter_b": _parse_parameter_billions(parameter) or _parse_parameter_billions(name),
        })
    return profiles


def fetch_target(target):
    first = None
    for url in preferred_urls(target["links"], target["note"]):
        response = fetch_response(url)
        if response is None:
            continue
        text = response_text(response[1], response[2])
        direct = direct_gpu_lines(text)
        profiles = parse_models(text)
        candidate = dict(target, url=url, response=response, text=text, direct_lines=direct, profiles=profiles)
        if direct or profiles:
            return candidate
        first = first or candidate
    return first or dict(target, url="", response=None, text="", direct_lines=[], profiles=[])


def decision(item):
    if item["response"] is None:
        return "unavailable", "当前未取得可复核响应", ""
    if item["direct_lines"]:
        if item["level"] == "高":
            return "keep_high", "当前响应含直接GPU字段", ""
        return "promote_high", "当前响应含直接GPU字段", ""
    profiles = item["profiles"]
    if not profiles:
        return "needs_review", "未从当前响应提取到可用于v7候选判定的模型列表", ""
    cloud = [profile for profile in profiles if profile["origin"] == "cloud"]
    local = [profile for profile in profiles if profile["origin"] == "local"]
    if not local:
        return "should_low_cloud", "模型均为明确cloud/远程来源，v7会从本地GPU候选中排除", "、".join(profile["name"] for profile in cloud[:4])
    retained = []
    excluded = []
    for profile in local:
        if profile["kind"] in {"embedding", "reranker"}:
            excluded.append("%s(%s)" % (profile["name"], profile["kind"]))
        elif profile["kind"] == "generative" and profile["parameter_b"] is not None and profile["parameter_b"] < 10:
            excluded.append("%s(%.3gB<10B)" % (profile["name"], profile["parameter_b"]))
        else:
            retained.append(profile)
    if not retained:
        return "should_low_filtered", "本地模型均属于v7排除类型", "、".join(excluded[:4])
    names = "、".join(profile["name"] for profile in retained[:4])
    if item["level"] == "高":
        return "needs_review_high", "当前无直接GPU字段，仅有可保留本地模型候选；v7候选逻辑只支持中等", names
    return "keep_medium", "存在可保留的本地GPU候选，未发现直接GPU字段", names


def main():
    parser = argparse.ArgumentParser(description="被动审计当前GPU候选是否符合v7模型规则")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--workers", type=int, default=8)
    args = parser.parse_args()

    workbook = load_workbook(args.input)
    targets = collect_targets(workbook)
    rows = []
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = [pool.submit(fetch_target, target) for target in targets]
        for future in as_completed(futures):
            item = future.result()
            status, reason, detail = decision(item)
            rows.append({
                "sheet": item["sheet"], "row": item["row"], "ip": item["ip"], "port": item["port"],
                "current_likelihood": item["level"], "status": status, "reason": reason,
                "detail": detail, "url": item["url"], "models": item["profiles"],
            })
    rows.sort(key=lambda item: (item["status"], item["sheet"], item["row"]))
    summary = {status: sum(item["status"] == status for item in rows) for status in sorted({item["status"] for item in rows})}
    args.report.write_text(json.dumps({"total": len(rows), "summary": summary, "rows": rows}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"total": len(rows), "summary": summary, "report": str(args.report)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
