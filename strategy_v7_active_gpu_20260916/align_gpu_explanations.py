#!/usr/bin/env python3
"""Align matched high-GPU explanations and screenshots to the live response."""
from __future__ import annotations

import argparse
import json
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook

from screenshot_evidence import (
    GPU_EVIDENCE_RE,
    add_metadata,
    attach_image,
    capture_response_card,
    detect_font,
    fetch_response,
    preferred_urls,
    remove_image_at,
    response_text,
)

GENERIC_MARKERS = {
    "gpu", "vram", "hbm", "cuda", "nvidia", "dcgm", "nvml", "accelerator",
}
NUMBER = r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?"


def headers(sheet):
    return {
        str(cell.value or "").strip(): index
        for index, cell in enumerate(sheet[1], start=1)
    }


def specific_markers(note):
    return {
        marker.lower()
        for marker in re.findall(r"[A-Za-z_][A-Za-z0-9_]*", str(note or ""))
        if GPU_EVIDENCE_RE.search(marker) and marker.lower() not in GENERIC_MARKERS
    }


def matching_lines(text, markers):
    return [
        line.strip()
        for line in text.splitlines()
        if any(marker in line.lower() for marker in markers)
    ]


def json_vram_facts(content_type, body):
    text = response_text(content_type, body)
    try:
        value = json.loads(text)
    except json.JSONDecodeError:
        return []
    models = value.get("models", []) if isinstance(value, dict) else []
    facts = []
    for model in models if isinstance(models, list) else []:
        if not isinstance(model, dict):
            continue
        size_vram = model.get("size_vram")
        if isinstance(size_vram, (int, float)) and size_vram > 0:
            facts.append("size_vram=%s" % size_vram)
        if len(facts) >= 3:
            break
    return facts


def line_facts(lines):
    facts = []
    seen = set()

    def add(fact):
        if fact and fact not in seen:
            seen.add(fact)
            facts.append(fact)

    for line in lines:
        metric = re.match(r"\s*([A-Za-z_:][A-Za-z0-9_:]*)\s*(\{[^}]*\})?\s+(%s)\s*$" % NUMBER, line)
        if metric and GPU_EVIDENCE_RE.search(metric.group(1)):
            name, labels, value = metric.groups()
            label_bits = []
            if labels:
                for key in ("gpu", "device", "modelName", "modelname"):
                    match = re.search(r"\b%s=\"([^\"]+)\"" % re.escape(key), labels)
                    if match:
                        label_bits.append("%s=%s" % (key, match.group(1)))
            fact = "%s=%s" % (name, value)
            if label_bits:
                fact += "（%s）" % "，".join(label_bits[:3])
            add(fact)

        for key, value in re.findall(
            r"[\"']?([A-Za-z_][A-Za-z0-9_]*)[\"']?\s*[:=]\s*[\"']?([^\"',}\]\s]+)",
            line,
        ):
            if (
                GPU_EVIDENCE_RE.search(key)
                and key.lower() not in GENERIC_MARKERS
                and not key.lower().endswith("_override")
                and value.lower() not in {"none", "null", "false", "true"}
            ):
                add("%s=%s" % (key, value))
        if len(facts) >= 4:
            break
    return facts[:4]


def aligned_note(url, response, old_note, lines):
    status, content_type, body = response
    facts = json_vram_facts(content_type, body) if "/api/ps" in url else []
    if not facts:
        facts = line_facts(lines)
    if not facts:
        raise ValueError("matched response produced no concise GPU facts")
    path = urlsplit(url).path or "/"
    fact_text = "、".join(facts)
    if "/api/ps" in path or any("size_vram=" in fact for fact in facts):
        return (
            "%s响应中显示%s；这是本机GPU显存占用的直接证据，"
            "因此GPU算力可能性高。" % (path, fact_text)
        )
    return (
        "%s响应中显示%s；这些字段属于本机GPU指标证据，"
        "因此GPU算力可能性高。" % (path, fact_text)
    )


def fetch_target(target):
    for url in preferred_urls(target["links"], target["old_note"]):
        response = fetch_response(url)
        if response is None:
            continue
        text = response_text(response[1], response[2])
        markers = specific_markers(target["old_note"])
        lines = matching_lines(text, markers)
        if lines:
            target = dict(target)
            target.update({"url": url, "response": response, "lines": lines})
            return target
        return None
    return None


def collect_targets(workbook):
    targets = []
    for sheet in workbook.worksheets:
        columns = headers(sheet)
        required = {
            "IP地址", "端口号", "访问链接", "证明截图",
            "存在GPU算力概率", "存在GPU算力的说明",
        }
        if not required.issubset(columns):
            continue
        image_rows = {
            image.anchor._from.row + 1
            for image in sheet._images
            if hasattr(image.anchor, "_from")
            and image.anchor._from.col + 1 == columns["证明截图"]
        }
        for row in range(2, sheet.max_row + 1):
            if not sheet.cell(row, columns["IP地址"]).value:
                continue
            if "高" not in str(sheet.cell(row, columns["存在GPU算力概率"]).value or ""):
                continue
            if row not in image_rows:
                continue
            targets.append({
                "sheet": sheet.title,
                "row": row,
                "image_column": columns["证明截图"],
                "note_column": columns["存在GPU算力的说明"],
                "ip": str(sheet.cell(row, columns["IP地址"]).value),
                "port": str(sheet.cell(row, columns["端口号"]).value or ""),
                "links": str(sheet.cell(row, columns["访问链接"]).value or ""),
                "old_note": str(sheet.cell(row, columns["存在GPU算力的说明"]).value or ""),
            })
    return targets


def main():
    parser = argparse.ArgumentParser(description="对齐GPU算力说明与证据截图")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--assets", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--expected", type=int, default=37)
    parser.add_argument("--workers", type=int, default=8)
    args = parser.parse_args()

    workbook = load_workbook(args.input)
    candidates = collect_targets(workbook)
    matched = []
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = [pool.submit(fetch_target, target) for target in candidates]
        for future in as_completed(futures):
            result = future.result()
            if result:
                matched.append(result)
    matched.sort(key=lambda item: (item["sheet"], item["row"]))
    if len(matched) != args.expected:
        raise RuntimeError("expected %d matched rows, got %d" % (args.expected, len(matched)))

    args.assets.mkdir(parents=True, exist_ok=True)
    font_path = str(detect_font())
    report = []
    for item in matched:
        sheet = workbook[item["sheet"]]
        note = aligned_note(item["url"], item["response"], item["old_note"], item["lines"])
        image_path = args.assets / ("%s_%s.png" % (item["sheet"], item["row"]))
        capture_response_card(item["url"], item["response"], image_path, font_path, note)
        add_metadata(image_path, item["url"], font_path)
        remove_image_at(sheet, item["row"], item["image_column"])
        attach_image(sheet, image_path, item["row"], item["image_column"])
        sheet.cell(item["row"], item["note_column"]).value = note
        report.append({
            "sheet": item["sheet"], "row": item["row"],
            "ip": item["ip"], "port": item["port"],
            "old_note": item["old_note"], "new_note": note,
            "url": item["url"],
        })

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(args.output)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print("aligned=%d output=%s" % (len(report), args.output))


if __name__ == "__main__":
    main()
