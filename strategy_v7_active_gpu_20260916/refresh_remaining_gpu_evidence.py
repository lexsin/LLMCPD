#!/usr/bin/env python3
"""Refresh evidence cards for unaligned high-GPU and all medium-GPU records.

The 37 high-risk rows already aligned by ``align_gpu_explanations.py`` are
excluded through its JSON report.  Every refreshed card is generated from the
selected live HTTP response, so a review does not depend on the browser's
initial viewport.
"""
from __future__ import annotations

import argparse
import json
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook

from align_gpu_explanations import GENERIC_MARKERS, line_facts
from screenshot_evidence import (
    GPU_EVIDENCE_RE,
    add_metadata,
    attach_image,
    capture_response_card,
    detect_font,
    fetch_response,
    preferred_urls,
    remove_image_at,
    response_text,
)


REQUIRED_HEADERS = {
    "IP地址", "端口号", "访问链接", "证明截图",
    "存在GPU算力概率", "存在GPU算力的说明",
}
DIRECT_GPU_TOKEN_RE = re.compile(
    r"gpu|vram|hbm|cuda|nvidia|dcgm|nvml|accelerator|size_vram",
    flags=re.IGNORECASE,
)


def headers(sheet):
    return {
        str(cell.value or "").strip(): column
        for column, cell in enumerate(sheet[1], start=1)
    }


def image_rows(sheet, image_column):
    return {
        image.anchor._from.row + 1
        for image in sheet._images
        if hasattr(image.anchor, "_from")
        and image.anchor._from.col + 1 == image_column
    }


def excluded_rows(report_path):
    value = json.loads(report_path.read_text(encoding="utf-8"))
    if not isinstance(value, list):
        raise ValueError("aligned report must be a list")
    return {(str(item["sheet"]), int(item["row"])) for item in value}


def direct_gpu_lines(text):
    """Return concise lines carrying a concrete, non-placeholder GPU fact."""
    lines = []
    for line in text.splitlines():
        # ``memory_total`` alone can describe CPU/RAM metrics.  Only fields
        # that themselves name GPU/VRAM/CUDA/NVIDIA-class hardware qualify.
        if not DIRECT_GPU_TOKEN_RE.search(line):
            continue
        if line_facts([line]):
            lines.append(line.strip())
    return lines


def select_live_response(target):
    for url in preferred_urls(target["links"], target["old_note"]):
        response = fetch_response(url)
        if response is None:
            continue
        text = response_text(response[1], response[2])
        direct_lines = direct_gpu_lines(text)
        chosen = dict(target)
        chosen.update({
            "url": url,
            "response": response,
            "text": text,
            "direct_lines": direct_lines,
        })
        return chosen
    return dict(target, url="", response=None, text="", direct_lines=[])


def direct_note(url, response, lines, probability):
    facts = line_facts(lines)
    if not facts:
        raise ValueError("direct GPU lines did not yield facts")
    path = urlsplit(url).path or "/"
    if "/api/ps" in path or any("size_vram=" in fact for fact in facts):
        evidence = "这是本机GPU显存占用的直接证据"
    else:
        evidence = "这些字段属于本机GPU指标证据"
    conclusion = "因此GPU算力可能性高。" if probability == "高" else "截图可作为GPU算力佐证。"
    return "%s响应中显示%s；%s，%s" % (
        path, "、".join(facts), evidence, conclusion,
    )


def model_identifiers(text):
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        parsed = None
    candidates = []

    def visit(value):
        if isinstance(value, dict):
            for key in ("id", "name", "model"):
                item = value.get(key)
                if isinstance(item, str) and item.strip():
                    candidates.append(item.strip())
            for item in value.values():
                visit(item)
        elif isinstance(value, list):
            for item in value[:12]:
                visit(item)

    if parsed is not None:
        visit(parsed)
    seen = set()
    result = []
    for item in candidates:
        if item not in seen:
            result.append(item)
            seen.add(item)
        if len(result) >= 4:
            break
    return result


def non_direct_note(url, text, probability):
    path = urlsplit(url).path or "/"
    models = model_identifiers(text)
    model_part = "，返回模型标识%s" % "、".join(models) if models else ""
    if probability == "中":
        return (
            "%s响应可访问%s；当前响应未出现GPU设备、显存占用或运行指标，"
            "因此GPU算力可能性中等。" % (path, model_part)
        )
    return (
        "%s响应可访问%s；当前响应未出现GPU设备、显存占用或运行指标，"
        "当前截图不构成直接GPU证据。" % (path, model_part)
    )


def unavailable_note():
    return "本次未获取到可用于复核的HTTP响应，未保留截图。"


def risk_level(value):
    """Normalize the workbook's verbose likelihood labels to 高 / 中."""
    text = str(value or "").strip()
    if "高" in text:
        return "高"
    if "中" in text:
        return "中"
    return ""


def collect_targets(workbook, excluded):
    targets = []
    for sheet in workbook.worksheets:
        columns = headers(sheet)
        if not REQUIRED_HEADERS.issubset(columns):
            continue
        for row in range(2, sheet.max_row + 1):
            if not sheet.cell(row, columns["IP地址"]).value:
                continue
            probability = risk_level(sheet.cell(row, columns["存在GPU算力概率"]).value)
            if probability not in {"高", "中"}:
                continue
            if probability == "高" and (sheet.title, row) in excluded:
                continue
            targets.append({
                "sheet": sheet.title,
                "row": row,
                "probability": probability,
                "image_column": columns["证明截图"],
                "note_column": columns["存在GPU算力的说明"],
                "ip": str(sheet.cell(row, columns["IP地址"]).value),
                "port": str(sheet.cell(row, columns["端口号"]).value or ""),
                "links": str(sheet.cell(row, columns["访问链接"]).value or ""),
                "old_note": str(sheet.cell(row, columns["存在GPU算力的说明"]).value or ""),
            })
    return targets


def main():
    parser = argparse.ArgumentParser(description="重跑未对齐高风险与全部中风险GPU证据")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--assets", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--exclude-report", type=Path, required=True)
    parser.add_argument("--workers", type=int, default=8)
    args = parser.parse_args()

    workbook = load_workbook(args.input)
    excluded = excluded_rows(args.exclude_report)
    targets = collect_targets(workbook, excluded)
    expected = 165
    if len(targets) != expected:
        raise RuntimeError("expected %d targets, got %d" % (expected, len(targets)))

    results = []
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = [pool.submit(select_live_response, item) for item in targets]
        for future in as_completed(futures):
            results.append(future.result())
    results.sort(key=lambda item: (item["sheet"], item["row"]))

    args.assets.mkdir(parents=True, exist_ok=True)
    font_path = str(detect_font())
    report = []
    for item in results:
        sheet = workbook[item["sheet"]]
        current_images = image_rows(sheet, item["image_column"])
        if item["response"] is None:
            if item["row"] in current_images:
                remove_image_at(sheet, item["row"], item["image_column"])
            note = unavailable_note()
            status = "unreachable"
            url = ""
        else:
            if item["direct_lines"]:
                note = direct_note(
                    item["url"], item["response"], item["direct_lines"], item["probability"],
                )
                status = "direct_evidence"
            else:
                note = non_direct_note(item["url"], item["text"], item["probability"])
                status = "service_only"
            image_path = args.assets / ("%s_%s.png" % (item["sheet"], item["row"]))
            capture_response_card(item["url"], item["response"], image_path, font_path, note)
            add_metadata(image_path, item["url"], font_path)
            if item["row"] in current_images:
                remove_image_at(sheet, item["row"], item["image_column"])
            attach_image(sheet, image_path, item["row"], item["image_column"])
            url = item["url"]
        sheet.cell(item["row"], item["note_column"]).value = note
        report.append({
            "sheet": item["sheet"], "row": item["row"], "ip": item["ip"],
            "port": item["port"], "probability": item["probability"],
            "status": status, "url": url, "old_note": item["old_note"], "new_note": note,
        })

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(args.output)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    counts = {key: sum(item["status"] == key for item in report) for key in ("direct_evidence", "service_only", "unreachable")}
    print("targets=%d direct=%d service_only=%d unreachable=%d output=%s" % (
        len(report), counts["direct_evidence"], counts["service_only"], counts["unreachable"], args.output,
    ))


if __name__ == "__main__":
    main()
