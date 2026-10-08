#!/usr/bin/env python3
"""Normalize carrier enterprise labels in SQLite locations and result workbooks."""

import argparse
import re
import shutil
import sqlite3
from pathlib import Path

from openpyxl import load_workbook


def normalize(value: str) -> str:
    text = str(value or "").strip()
    text = re.sub(r"^[0-9]+", "", text)
    carrier_match = re.match(r"^(?:中国)?(移动|电信|联通)(.+)$", text)
    if carrier_match:
        carrier, province = carrier_match.groups()
    else:
        carrier_match = re.match(r"^(.+?)(?:中国)?(移动|电信|联通)$", text)
        if not carrier_match:
            return text
        province, carrier = carrier_match.groups()
    province = province.strip()
    if province == "内蒙":
        province = "内蒙古"
    return province + carrier


def workbook(source: Path, output: Path) -> None:
    book = load_workbook(source, data_only=False)
    renamed = {}
    for sheet in book.worksheets:
        if sheet.title == "汇总统计":
            continue
        clean = normalize(sheet.title)
        if clean != sheet.title:
            if clean in book.sheetnames:
                raise ValueError("sheet name collision: %s -> %s" % (sheet.title, clean))
            renamed[sheet.title] = clean
            sheet.title = clean
    if "汇总统计" in book.sheetnames:
        summary = book["汇总统计"]
        for row in range(2, summary.max_row + 1):
            value = summary.cell(row, 1).value
            if value and str(value) != "合计":
                summary.cell(row, 1).value = normalize(str(value))
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    book.close()
    print("workbook_renamed=%d output=%s" % (len(renamed), output))


def sqlite_locations(database: Path, backup: Path) -> None:
    shutil.copy2(database, backup)
    connection = sqlite3.connect(database)
    try:
        tables = ("source_ip_ranges", "source_ip_ranges_batch2")
        changes = []
        with connection:
            for table in tables:
                values = [row[0] for row in connection.execute("select distinct location from " + table)]
                for old in values:
                    new = normalize(old)
                    if new != old:
                        connection.execute("update " + table + " set location=? where location=?", (new, old))
                        changes.append((table, old, new))
    finally:
        connection.close()
    print("sqlite_renamed=%d backup=%s" % (len(changes), backup))
    for table, old, new in changes:
        print("%s: %s -> %s" % (table, old, new))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--database", type=Path)
    parser.add_argument("--backup", type=Path)
    args = parser.parse_args()
    if args.workbook:
        if not args.output:
            parser.error("--output is required with --workbook")
        workbook(args.workbook, args.output)
    if args.database:
        if not args.backup:
            parser.error("--backup is required with --database")
        sqlite_locations(args.database, args.backup)


if __name__ == "__main__":
    main()
