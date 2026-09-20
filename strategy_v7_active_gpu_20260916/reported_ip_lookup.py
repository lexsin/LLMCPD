# -*- coding: utf-8 -*-
"""Mark workbook rows whose IP belongs to an imported reporting range."""

import bisect
import ipaddress
import sqlite3
from copy import copy
from pathlib import Path

from openpyxl.utils import get_column_letter


def load_reported_intervals(database: Path, table: str):
    uri = f"file:{database.resolve().as_posix()}?mode=ro"
    with sqlite3.connect(uri, uri=True) as connection:
        exists = connection.execute(
            "SELECT 1 FROM sqlite_master "
            "WHERE type IN ('table', 'view') AND name = ?",
            (table,),
        ).fetchone()
        if not exists:
            raise ValueError(f"SQLite table not found: {table}")
        quoted_table = '"' + table.replace('"', '""') + '"'
        rows = connection.execute(
            f"SELECT start_ip_int, end_ip_int FROM {quoted_table} "
            "WHERE start_ip_int IS NOT NULL AND end_ip_int IS NOT NULL "
            "ORDER BY start_ip_int, end_ip_int"
        )
        merged = []
        for start, end in rows:
            start, end = int(start), int(end)
            if end < start:
                start, end = end, start
            if not merged or start > merged[-1][1] + 1:
                merged.append([start, end])
            elif end > merged[-1][1]:
                merged[-1][1] = end
    starts = [start for start, _ in merged]
    ends = [end for _, end in merged]

    def contains(value):
        try:
            address = ipaddress.ip_address(str(value or "").strip())
        except ValueError:
            return False
        if address.version != 4:
            return False
        number = int(address)
        index = bisect.bisect_right(starts, number) - 1
        return index >= 0 and number <= ends[index]

    return contains, len(merged)


def add_reported_column(sheet, contains, header="是否上报"):
    header_row = None
    ip_column = None
    output_column = None
    for row in range(1, min(sheet.max_row, 10) + 1):
        for column in range(1, sheet.max_column + 1):
            value = str(sheet.cell(row, column).value or "").strip()
            if value == "IP地址":
                header_row, ip_column = row, column
            if value == header:
                output_column = column
        if header_row and ip_column:
            break
    if header_row is None or ip_column is None:
        return 0, 0
    output_column = output_column or sheet.max_column + 1
    header_cell = sheet.cell(header_row, output_column, header)
    style_source = sheet.cell(header_row, max(1, output_column - 1))
    if style_source.has_style:
        header_cell._style = copy(style_source._style)
    header_cell.alignment = copy(style_source.alignment)
    sheet.column_dimensions[get_column_letter(output_column)].width = 12

    yes = no = 0
    for row in range(header_row + 1, sheet.max_row + 1):
        ip = sheet.cell(row, ip_column).value
        if ip in (None, ""):
            continue
        reported = contains(ip)
        target = sheet.cell(row, output_column, "是" if reported else "否")
        body_source = sheet.cell(row, max(1, output_column - 1))
        if body_source.has_style:
            target._style = copy(body_source._style)
        target.alignment = copy(body_source.alignment)
        yes += int(reported)
        no += int(not reported)
    return yes, no
