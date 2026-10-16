import csv, json
from pathlib import Path
from openpyxl import load_workbook

base_dir = Path('data')
stems = sorted([p.stem for p in base_dir.glob('3*.xlsx') if not p.name.endswith('_最终版.xlsx')])

def count_csv(path):
    if not path.exists():
        return None
    for enc in ('utf-8-sig','utf-8','gbk'):
        try:
            with path.open(encoding=enc, newline='') as f:
                return sum(1 for _ in csv.DictReader(f))
        except UnicodeDecodeError:
            continue
    with path.open(newline='') as f:
        return sum(1 for _ in csv.DictReader(f))

def scan_stats(path):
    if not path.exists():
        return {'scan_rows': None, 'open_ip_rows': None, 'open_port_items': None}
    rows = []
    for enc in ('utf-8-sig','utf-8','gbk'):
        try:
            with path.open(encoding=enc, newline='') as f:
                rows = list(csv.DictReader(f))
            break
        except UnicodeDecodeError:
            continue
    open_ip_rows = 0
    open_port_items = 0
    for r in rows:
        ports = (r.get('open_ports') or '').strip()
        if ports:
            parts = [x.strip() for x in ports.replace(',', ';').split(';') if x.strip()]
            if parts:
                open_ip_rows += 1
                open_port_items += len(set(parts))
    return {'scan_rows': len(rows), 'open_ip_rows': open_ip_rows, 'open_port_items': open_port_items}

def final_stats(path):
    out = {'gpu_high': None, 'gpu_medium': None, 'gpu_low': None, 'empty_sheet': None}
    if not path.exists():
        return out
    wb = load_workbook(path, read_only=True, data_only=True)
    def rows(sheet):
        if sheet not in wb.sheetnames:
            return 0
        ws = wb[sheet]
        return max(ws.max_row - 1, 0)
    out['gpu_high'] = rows('GPU可能性高')
    out['gpu_medium'] = rows('GPU可能性中')
    out['gpu_low'] = rows('GPU可能性低')
    out['empty_sheet'] = rows('无GPU可能性结果') if '无GPU可能性结果' in wb.sheetnames else 0
    return out

records = []
for stem in stems:
    rec = {'stem': stem}
    rec['ip_rows'] = count_csv(base_dir / f'{stem}_result.csv')
    rec.update(scan_stats(base_dir / f'{stem}_scan_result.csv'))
    rec['scan_ports_rows'] = count_csv(base_dir / f'{stem}_scan_ports.csv')
    rec['llm_rows'] = count_csv(base_dir / f'{stem}_llm.csv')
    rec.update(final_stats(base_dir / f'{stem}_llm_最终版.xlsx'))
    records.append(rec)

totals = {
    'files': len(records),
    'final_files': sum(1 for r in records if (base_dir / (r['stem'] + '_llm_最终版.xlsx')).exists()),
    'ip_rows': sum(r['ip_rows'] or 0 for r in records),
    'scan_rows': sum(r['scan_rows'] or 0 for r in records),
    'open_ip_rows': sum(r['open_ip_rows'] or 0 for r in records),
    'open_port_items': sum(r['open_port_items'] or 0 for r in records),
    'scan_ports_rows': sum(r['scan_ports_rows'] or 0 for r in records),
    'llm_rows': sum(r['llm_rows'] or 0 for r in records),
    'gpu_high': sum(r['gpu_high'] or 0 for r in records),
    'gpu_medium': sum(r['gpu_medium'] or 0 for r in records),
    'gpu_low': sum(r['gpu_low'] or 0 for r in records),
}
print(json.dumps({'totals': totals, 'records': records}, ensure_ascii=False, indent=2))
