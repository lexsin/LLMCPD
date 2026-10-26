import importlib.util
import json
import os
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

from llm_detect.v21.common import ContractError, atomic_json, beneath, proof_hash
from llm_detect.v21.config import Config, DEFAULT_RELEASE, SERVICE_ROOT
from llm_detect.v21.releases import Releases, frozen_rules
from llm_detect.v21.runtime import Runtime
from llm_detect.v21.stage_worker import write_csv
from .test_contract import prepared


def test_safe_paths(tmp_path):
    for path in ("../other","/etc/passwd","C:/temp/file","tasks/../../outside"):
        with pytest.raises(ContractError):beneath(tmp_path,path)
    assert beneath(tmp_path,"tasks/1/runs/2").is_relative_to(tmp_path)


def test_release_integrity_and_no_fallback(tmp_path):
    releases=Releases(SERVICE_ROOT/"vendor/v7",DEFAULT_RELEASE)
    release=releases.get()
    assert release.version.startswith(DEFAULT_RELEASE+"+sha256:")
    assert releases.get(release.version).version==release.version
    with pytest.raises(ContractError):releases.get("nonexistent-fixed-release+sha256:"+"0"*64)
    import shutil
    shutil.copytree(release.path,tmp_path/DEFAULT_RELEASE,ignore=shutil.ignore_patterns("__pycache__"))
    (tmp_path/DEFAULT_RELEASE/"scan_llm.py").write_text("# tampered",encoding="utf-8")
    with pytest.raises(ContractError,match="differs"):Releases(tmp_path,DEFAULT_RELEASE).get(release.version)


def test_rules_optional_only_before_first_start():
    release=Releases(SERVICE_ROOT/"vendor/v7",DEFAULT_RELEASE).get()
    snapshot,checksum=frozen_rules({},release)
    assert frozen_rules({"rule_snapshot_json":json.dumps(snapshot),"rule_hash":checksum,"started_at":"started"},release)==(snapshot,checksum)
    with pytest.raises(ContractError):frozen_rules({"started_at":"started"},release)
    with pytest.raises(ContractError):frozen_rules({"rule_snapshot_json":json.dumps(snapshot)},release)


@pytest.mark.parametrize("manifest", [{}, {"release_id": DEFAULT_RELEASE, "upstream": "fixture", "files": {}, "adapter_files": {}}])
def test_incomplete_release_manifest_has_explicit_error(tmp_path, manifest):
    directory=tmp_path/DEFAULT_RELEASE
    directory.mkdir()
    atomic_json(directory/"release.json",manifest)
    with pytest.raises(ContractError,match="incomplete|invalid SHA-256"):
        Releases(tmp_path,DEFAULT_RELEASE).get()


def test_remote_disconnect_does_not_start_replacement_or_complete(tmp_path,monkeypatch):
    db,repo,lease,sources=prepared(tmp_path)
    cfg=Config(data_root=tmp_path,control_seconds=0.001)
    runtime=Runtime(cfg,repo)
    fake=SimpleNamespace(root=beneath(tmp_path,sources[0]["workspace_relative_path"]))
    calls=[]
    def prepare():
        calls.append("prepare-same-job")
        raise OSError("SSH unavailable")
    fake.prepare=prepare
    runtime.remote_jobs=[fake]
    checked=[]
    def control(*args):
        checked.append(1)
        return len(checked)>2
    monkeypatch.setattr(runtime,"control",control)
    assert runtime.remote_stage(lease,None,sources[0],{},tmp_path) is None
    assert calls==["prepare-same-job","prepare-same-job"]
    assert repo.get_run(lease.run_id)["status"]=="RUNNING"


def test_remote_agent_rejects_stale_owner_and_binding(tmp_path,monkeypatch):
    # Agent is Linux-only; mock flock for pure RPC ownership tests on Windows.
    if os.name!="posix":
        monkeypatch.setitem(sys.modules,"fcntl",SimpleNamespace(LOCK_EX=1,flock=lambda *a:None))
    spec=importlib.util.spec_from_file_location("offline_remote_guardian",SERVICE_ROOT/"vendor/v7"/DEFAULT_RELEASE/"service_remote.py")
    agent=importlib.util.module_from_spec(spec);spec.loader.exec_module(agent)
    if os.name!="posix":
        # Windows cannot open a directory for fsync; Linux runs exercise the real save().
        monkeypatch.setattr(agent,"save",atomic_json)
    atomic_json(tmp_path/"owner.json",{"epoch":3,"binding":"fixed","token":"old"})
    with pytest.raises(RuntimeError,match="stale"):agent.rpc("poll",tmp_path,{"epoch":2,"binding":"fixed","ttl":60})
    with pytest.raises(RuntimeError,match="binding"):agent.rpc("poll",tmp_path,{"epoch":4,"binding":"other","ttl":60})
    monkeypatch.setattr(agent,"members",lambda *a:[])
    assert agent.rpc("stop",tmp_path,{"epoch":4,"binding":"fixed","ttl":60})["state"]=="stopped"


def test_run_workbook_keeps_sources_separate(tmp_path,monkeypatch):
    release=SERVICE_ROOT/"vendor/v7"/DEFAULT_RELEASE
    monkeypatch.syspath_prepend(str(release))
    import finalize_llm_results
    monkeypatch.setattr(finalize_llm_results,"fill_screenshots",lambda *a,**kw:0)
    from llm_detect.v21.reports import build
    sources=[]
    for sid in (11,12):
        root=tmp_path/str(sid);root.mkdir()
        write_csv(root/"committed.csv",[{"ip":"192.0.2.1","port":"80","protocol":"http","gpu_likelihood":"高","npu_likelihood":"未知"}])
        sources.append({"id":sid,"file_id":sid,"source_name":f"file-{sid}","status":"SUCCESS","root":str(root)})
    build({"release_path":str(release),"output":str(tmp_path/"run.xlsx"),"sources":sources})
    from openpyxl import load_workbook
    workbook=load_workbook(tmp_path/"run.xlsx")
    assert workbook.sheetnames[0]=="来源清单"
    assert any(n.startswith("S11_") for n in workbook.sheetnames)
    assert any(n.startswith("S12_") for n in workbook.sheetnames)
    workbook.close()


def test_empty_run_report_is_valid_workbook(tmp_path,monkeypatch):
    release=SERVICE_ROOT/"vendor/v7"/DEFAULT_RELEASE
    monkeypatch.syspath_prepend(str(release))
    import finalize_llm_results
    monkeypatch.setattr(finalize_llm_results,"fill_screenshots",lambda *a,**kw:0)
    from llm_detect.v21.reports import build
    root=tmp_path/"source";root.mkdir();write_csv(root/"committed.csv",[])
    build({"release_path":str(release),"output":str(tmp_path/"empty.xlsx"),"sources":[{"id":1,"file_id":1,"source_name":"empty","status":"SUCCESS","root":str(root)}]})
    assert (tmp_path/"empty.xlsx").stat().st_size>0
