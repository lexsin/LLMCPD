"""Offline RUN workbook assembly. Never merge evidence by IP across source groups."""
from __future__ import annotations

import copy
import io
import json
import sys
from pathlib import Path


def build(context):
    sys.path.insert(0, context["release_path"])
    import service_hooks  # keep the run fence across offline image-embedding children
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image
    import finalize_llm_results as finalizer
    from .report_evidence import prepare_manifest
    workbook = Workbook()
    index = workbook.active
    index.title = "来源清单"
    index.append(["来源组ID", "文件ID", "文件名称", "执行状态", "工作表"])
    for source in context["sources"]:
        root = Path(source["root"])
        temporary = root / "run_report_source.xlsx"
        data = finalizer.read_scan(root / "committed.csv")
        finalizer.write_result(data, temporary)
        code = finalizer.fill_screenshots(temporary, prepare_manifest(root), manifest_only=True)
        if code:
            raise RuntimeError("offline evidence embedding failed")
        part = load_workbook(temporary)
        for sheet in part.worksheets:
            title = (f"S{source['id']}_" + sheet.title)[:31]
            target = workbook.create_sheet(title)
            index.append([source["id"], source["file_id"], source["source_name"], source["status"], title])
            for row in sheet:
                for cell in row:
                    dest = target.cell(cell.row, cell.column, cell.value)
                    if cell.has_style:
                        dest._style = copy.copy(cell._style)
                    dest.number_format = cell.number_format
                    if cell.hyperlink:
                        dest._hyperlink = copy.copy(cell.hyperlink)
                    if cell.alignment:
                        dest.alignment = copy.copy(cell.alignment)
            for key, dimension in sheet.column_dimensions.items():
                target.column_dimensions[key] = copy.copy(dimension)
            for key, dimension in sheet.row_dimensions.items():
                target.row_dimensions[key] = copy.copy(dimension)
            for merged in sheet.merged_cells.ranges:
                target.merge_cells(str(merged))
            target.freeze_panes = sheet.freeze_panes
            target.auto_filter.ref = sheet.auto_filter.ref
            for image in sheet._images:
                replacement = Image(io.BytesIO(image._data()))
                replacement.width, replacement.height = image.width, image.height
                target.add_image(replacement, copy.deepcopy(image.anchor))
        part.close()
    workbook.save(context["output"])
    workbook.close()


if __name__ == "__main__":
    from .processes import start_lease_guard
    start_lease_guard()
    build(json.loads(Path(sys.argv[1]).read_text(encoding="utf-8")))
