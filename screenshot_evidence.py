#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fill Excel '证明截图' with evidence-aware screenshots (URL + Beijing timestamp)."""
from __future__ import annotations

import argparse
import shutil
import ssl
from datetime import datetime
from pathlib import Path
from subprocess import PIPE, STDOUT, TimeoutExpired, run
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

CONTEXT = ssl._create_unverified_context()
IMAGE_DISPLAY = (400, 180)
IMAGE_ROW_HEIGHT = 140
IMAGE_COL_WIDTH = 62
CHROME_CANDIDATES = [
    Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
    Path(r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
    Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
    Path("/usr/bin/google-chrome"),
    Path("/usr/bin/google-chrome-stable"),
    Path("/opt/google/chrome/chrome"),
    Path("/usr/bin/chromium-browser"),
    Path("/usr/bin/chromium"),
]
FONT_CANDIDATES = [
    Path(r"C:\Windows\Fonts\msyh.ttc"),
    Path(r"C:\Windows\Fonts\msyh.ttf"),
    Path("/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc"),
    Path("/usr/share/fonts/google-noto-sans-cjk-fonts/NotoSansCJK-Regular.ttc"),
    Path("/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc"),
    Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    Path("/usr/share/fonts/wqy-zenhei/wqy-zenhei.ttc"),
    Path("/usr/share/fonts/wqy-microhei/wqy-microhei.ttc"),
]


def find_column(sheet, header):
    for column in range(1, sheet.max_column + 1):
        if str(sheet.cell(1, column).value or "").strip() == header:
            return column
    return None


def detect_chrome():
    for name in ("google-chrome", "google-chrome-stable", "chromium-browser", "chromium", "chrome"):
        found = shutil.which(name)
        if found:
            return Path(found)
    for path in CHROME_CANDIDATES:
        if path.is_file():
            return path
    raise FileNotFoundError("未找到 Chrome/Chromium，无法截图")


def detect_font():
    for path in FONT_CANDIDATES:
        if path.is_file():
            return path
    raise FileNotFoundError("未找到中文字体，无法写入 URL/时间戳")


def split_urls(value):
    return [url.strip() for url in str(value or "").split("|||") if url.strip()]


def evidence_kind(note):
    text = str(note or "")
    lowered = text.lower()
    if "/api/ps" in lowered or "size_vram" in lowered or "显存" in text:
        return "vram"
    if any(token in lowered for token in ("/metrics", "dcgm", "gpu_memory", "vllm_gpu", "gpu_uuid", "nv_gpu")):
        return "metrics"
    return "models"


def score_url(url, kind):
    if kind == "vram":
        if "/api/ps" in url:
            return 0
        if "/v1/models" in url:
            return 2
        if "/api/tags" in url:
            return 3
        return 9
    if kind == "metrics":
        if "/metrics" in url:
            return 0
        if "/v1/models" in url:
            return 2
        return 9
    if "/v1/models" in url:
        return 0
    if "/api/tags" in url:
        return 1
    if "/info" in url:
        return 2
    if "/props" in url:
        return 3
    if "/api/version" in url:
        return 4
    return 8


def preferred_urls(links, note):
    kind = evidence_kind(note)
    urls = split_urls(links)
    return sorted(urls, key=lambda url: score_url(url, kind))


def reachable(url):
    try:
        with urlopen(Request(url, headers={"User-Agent": "Mozilla/5.0"}), timeout=8, context=CONTEXT) as response:
            body = response.read(256)
            return 200 <= response.status < 400 and bool(body)
    except Exception:
        return False


def capture(chrome, url, path):
    target = Path(path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    command = [
        str(chrome), "--headless", "--disable-gpu", "--no-first-run",
        "--no-sandbox", "--disable-dev-shm-usage",
        "--ignore-certificate-errors", "--virtual-time-budget=4000",
        "--window-size=1100,500", f"--screenshot={target}", url,
    ]
    try:
        result = run(command, stdout=PIPE, stderr=STDOUT, text=True, timeout=20)
    except TimeoutExpired:
        return False
    return result.returncode == 0 and target.exists() and target.stat().st_size > 0


def add_metadata(path, url, font_path):
    page = Image.open(path).convert("RGB")
    header_height, footer_height = 42, 36
    image = Image.new("RGB", (page.width, page.height + header_height + footer_height), "white")
    image.paste(page, (0, header_height))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, page.width, header_height), fill=(245, 245, 245))
    draw.line((0, header_height - 1, page.width, header_height - 1), fill=(210, 210, 210))
    draw.text((16, 10), url, font=ImageFont.truetype(font_path, 18), fill=(35, 35, 35))
    timestamp = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M:%S")
    draw.text(
        (16, page.height + header_height + 8),
        f"截图时间：{timestamp}（北京时间）",
        font=ImageFont.truetype(font_path, 16),
        fill=(75, 75, 75),
    )
    image.save(path)


def image_cells(sheet, image_column):
    cells = set()
    for image in sheet._images:
        anchor = image.anchor
        if hasattr(anchor, "_from") and anchor._from.col + 1 == image_column:
            cells.add(anchor._from.row + 1)
    return cells


def attach_image(sheet, path, row, column):
    """Bind screenshot to the 证明截图 cell (TwoCellAnchor, one cell).

    This is Excel 'Place over cells' sized to that cell, not 365 'Place in Cell'.
    WPS and older Excel can display it; the picture moves with the row.
    """
    image = ExcelImage(path)
    image.width, image.height = IMAGE_DISPLAY
    image.anchor = TwoCellAnchor(
        editAs="oneCell",
        _from=AnchorMarker(col=column - 1, colOff=0, row=row - 1, rowOff=0),
        to=AnchorMarker(col=column, colOff=0, row=row, rowOff=0),
    )
    sheet.add_image(image)
    sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 0, IMAGE_ROW_HEIGHT)
    letter = get_column_letter(column)
    if (sheet.column_dimensions[letter].width or 0) < 58:
        sheet.column_dimensions[letter].width = IMAGE_COL_WIDTH


def fill_workbook(source, output, assets, blank_only, sheet_name=None, limit=None):
    chrome = detect_chrome()
    font_path = str(detect_font())
    assets.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source)
    records = []
    filled = 0

    sheets = [workbook[sheet_name]] if sheet_name else workbook.worksheets
    for sheet in sheets:
        link_column = find_column(sheet, "访问链接")
        note_column = find_column(sheet, "存在GPU算力的说明")
        image_column = find_column(sheet, "证明截图")
        if not link_column or not image_column:
            continue
        occupied = image_cells(sheet, image_column)
        for row in range(2, sheet.max_row + 1):
            if limit is not None and filled >= limit:
                break
            if not sheet.cell(row, 1).value:
                continue
            if blank_only and row in occupied:
                continue
            urls = preferred_urls(
                sheet.cell(row, link_column).value,
                sheet.cell(row, note_column).value if note_column else "",
            )
            selected = next((url for url in urls if reachable(url)), None)
            if not selected:
                records.append((sheet.title, row, "未获取到有效响应", ""))
                continue
            path = assets / f"{sheet.title}_{row}.png"
            if not capture(chrome, selected, path):
                records.append((sheet.title, row, "截图生成失败", selected))
                continue
            add_metadata(path, selected, font_path)
            if row in occupied:
                sheet._images = [
                    image for image in sheet._images
                    if not (
                        hasattr(image.anchor, "_from")
                        and image.anchor._from.row + 1 == row
                        and image.anchor._from.col + 1 == image_column
                    )
                ]
            attach_image(sheet, path, row, image_column)
            records.append((sheet.title, row, "已填充", selected))
            filled += 1

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output)
    return records


def main():
    parser = argparse.ArgumentParser(description="为 LLM 结果 Excel 填充证明截图")
    parser.add_argument("command", choices=["fill", "retry-blank"])
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--assets", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--limit", type=int)
    args = parser.parse_args()

    source = args.input
    output = args.output or source
    assets = args.assets or (source.parent / "screenshot_assets")
    records = fill_workbook(
        source,
        output,
        assets,
        blank_only=(args.command == "retry-blank"),
        sheet_name=args.sheet,
        limit=args.limit,
    )
    for sheet, row, status, url in records:
        print(f"{sheet}!{row}: {status} {url}")
    print(f"输出文件: {output}")


if __name__ == "__main__":
    main()
