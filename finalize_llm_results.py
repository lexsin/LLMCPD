#!/usr/bin/env python3
import argparse
import re
import subprocess
import sys
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT_COLUMNS = [
    "IP地址", "端口号", "协议类型", "服务类型", "model_domain", "模型部署工具", "工具版本", "模型信息",
    "访问链接", "探测说明", "证明截图", "存在GPU算力概率", "存在GPU算力的说明", "是否为上报IP",
]

GPU_LABELS = {
    "高": "GPU算力可能性高", "中": "GPU算力可能性中等", "低": "GPU算力可能性较低", "未知": "GPU算力可能性未知",
    "high": "GPU算力可能性高", "medium": "GPU算力可能性中等", "low": "GPU算力可能性较低", "unknown": "GPU算力可能性未知",
}

COLUMN_WIDTHS = {
    "A": 18.6272727272727,
    "B": 8.72727272727273,
    "C": 13.0,
    "D": 13.0,
    "E": 13.0,
    "F": 13.0,
    "G": 10.0,
    "H": 18.0,
    "I": 16.3727272727273,
    "J": 30.0,
    "K": 62.0,
    "L": 11.2545454545455,
    "M": 34.1272727272727,
    "N": 14.0,
}


def read_scan(path: Path) -> pd.DataFrame:
    if path.suffix.lower() in {".xlsx", ".xls"}:
        return pd.read_excel(path)
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return pd.read_csv(path, encoding=enc)
        except UnicodeDecodeError:
            continue
    return pd.read_csv(path)


def read_reported_single_ips(path: Path | None) -> set[str]:
    if not path:
        return set()
    df = read_scan(path)
    if not {"起始IP", "终止IP"}.issubset(df.columns):
        return set()
    starts = df["起始IP"].astype(str).str.strip()
    ends = df["终止IP"].astype(str).str.strip()
    return set(starts[starts == ends])


def clean(value) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "null"} else text


def gpu_label(value: str) -> str:
    value = clean(value)
    return GPU_LABELS.get(value, GPU_LABELS.get(value.lower(), "GPU算力可能性未知"))


def extract_model_candidates(model_info: str, evidence: str) -> list[str]:
    text = clean(model_info)
    candidates = []
    if text and text != "未知":
        candidates.extend([part.strip() for part in text.split(",") if part.strip()])
    ev = clean(evidence)
    for pat in (r'"id"\s*:\s*"([^"]+)"', r'"name"\s*:\s*"([^"]+)"', r'"model"\s*:\s*"([^"]+)"'):
        for match in re.findall(pat, ev):
            if match and match not in candidates:
                candidates.append(match)
            if len(candidates) >= 8:
                break
        if len(candidates) >= 8:
            break
    cleaned = []
    for item in candidates:
        item = re.sub(r"^/model/", "", item.strip())
        item = item.strip("/")
        if item and not item.lower().startswith(("modelperm-", "sha256:", "gpu-")) and item not in cleaned:
            cleaned.append(item)
    return cleaned[:8]


def short_models(model_info: str, evidence: str) -> str:
    candidates = extract_model_candidates(model_info, evidence)
    return "、".join(candidates) if candidates else "未知"


def tool_display(deploy_tool: str, evidence: str) -> str:
    tool = clean(deploy_tool).lower()
    ev = clean(evidence).lower()
    if "vllm" in tool or "owned_by\":\"vllm" in ev:
        return "vLLM/OpenAI兼容"
    if "sglang" in tool or "owned_by\":\"sglang" in ev:
        return "SGLang/OpenAI兼容"
    if "xinference" in tool or "owned_by\":\"xinference" in ev or "<title>xinference</title>" in ev:
        return "Xinference/OpenAI兼容"
    if "tgi" in tool or "text-generation-inference" in tool:
        return "TGI/OpenAI兼容"
    if "ollama" in tool or "ollama is running" in ev or "/api/tags" in ev:
        return "Ollama/OpenAI兼容"
    if "gradio" in tool:
        return "Gradio/AI前端"
    if "one-api" in tool:
        return "One-API模型网关"
    if "openai" in tool:
        return "OpenAI兼容"
    return clean(deploy_tool) or "相关"


def model_summary(model_info: str, evidence: str, limit: int = 4) -> str:
    candidates = extract_model_candidates(model_info, evidence)
    if not candidates:
        return ""
    shown = "、".join(candidates[:limit])
    return shown + ("等" if len(candidates) > limit else "")


def split_cloud_models(model_info: str, evidence: str) -> tuple[list[str], list[str]]:
    local_models = []
    cloud_models = []
    for model in extract_model_candidates(model_info, evidence):
        if ":cloud" in model.lower():
            cloud_models.append(model)
        else:
            local_models.append(model)
    return local_models, cloud_models


def format_vram_bytes(value: str) -> str:
    try:
        size = int(value)
    except (TypeError, ValueError):
        return clean(value)
    return f"{size / (1024 ** 3):.1f} GiB"


def gpu_note(row) -> str:
    evidence = clean(row.get("evidence", ""))
    detail = clean(row.get("gpu_probe_detail", ""))
    gpu_evidence = clean(row.get("gpu_evidence", ""))
    likelihood = clean(row.get("gpu_likelihood", "未知"))
    service_type = clean(row.get("service_type", ""))
    models = model_summary(row.get("model_info", ""), evidence)
    tool = tool_display(row.get("deploy_tool", ""), evidence)

    if likelihood in {"高", "high"}:
        same_ip = re.search(r"same_ip_gpu_port=([^;]+)", gpu_evidence, flags=re.I)
        if same_ip:
            direct = re.sub(r"^same_ip_gpu_port=[^;]+;?\s*", "", gpu_evidence, flags=re.I)
            return (
                f"同一IP的{same_ip.group(1)}端口探测到直接GPU证据{direct}；"
                "当前端口已确认是LLM模型服务，因此关联判断本机存在GPU算力，GPU算力可能性高。"
            )
        vram_hits = re.findall(r"size_vram=(\d+)(?:\s+model=([^;]+))?", gpu_evidence, flags=re.I)
        if vram_hits:
            parts = []
            for size, model in vram_hits[:3]:
                model_text = clean(model) or "模型"
                parts.append(f"{model_text}占用约{format_vram_bytes(size)}显存")
            return (
                "探测响应显示Ollama /api/ps返回正在运行的模型，"
                + "、".join(parts)
                + "；这是本机GPU显存占用的直接证据，因此GPU算力可能性高。"
            )
        metric_hits = []
        for item in gpu_evidence.split(";"):
            item = item.strip()
            if item and item not in metric_hits:
                metric_hits.append(item)
        if metric_hits:
            return (
                "探测响应命中GPU设备或监控指标："
                + "、".join(metric_hits[:5])
                + "；该类指标属于较强的本机GPU证据，因此GPU算力可能性高。"
            )
        return "探测响应显示GPU设备指标、显存占用或运行状态，属于较强的本机GPU证据，因此GPU算力可能性高。"

    if likelihood in {"中", "medium"}:
        if models:
            model_text = f"返回模型标识{models}；"
        else:
            model_text = "模型列表为空或未返回明确模型标识；"
        return (
            f"探测响应确认{tool}服务接口可访问，{model_text}"
            "已识别到本地推理框架，但未获得GPU设备、显存占用或运行指标，"
            "因此GPU算力可能性中等。"
        )

    if likelihood in {"低", "low"}:
        frontend_tools = {
            "gradio/ai前端", "one-api模型网关", "open-webui", "librechat", "ai前端"
        }
        if service_type == "AI前端服务" or tool.lower() in frontend_tools:
            return (
                "探测响应显示AI前端页面或模型聚合网关可访问，但未确认本机模型推理接口，"
                "也未发现GPU设备、显存占用或运行指标，因此GPU算力可能性较低。"
            )
        local_models, cloud_models = split_cloud_models(row.get("model_info", ""), evidence)
        if cloud_models and not local_models:
            cloud_text = "、".join(cloud_models[:4]) + ("等" if len(cloud_models) > 4 else "")
            return (
                f"探测响应确认{tool}接口可访问，但仅返回云端/远程模型{cloud_text}；"
                "远程模型不能证明本机存在GPU算力，因此GPU算力可能性较低。"
            )
        if cloud_models and local_models:
            local_text = "、".join(local_models[:3]) + ("等" if len(local_models) > 3 else "")
            cloud_text = "、".join(cloud_models[:2]) + ("等" if len(cloud_models) > 2 else "")
            return (
                f"探测响应确认{tool}接口可访问，模型列表同时包含本地标识{local_text}"
                f"和云端标识{cloud_text}；但未发现模型正在本机GPU运行或占用显存的证据，"
                "因此GPU算力可能性较低。"
            )
        if local_models:
            local_text = "、".join(local_models[:4]) + ("等" if len(local_models) > 4 else "")
            return (
                f"探测响应确认{tool}模型接口可访问，并返回模型标识{local_text}；"
                "但未发现GPU设备、显存占用或运行指标，现有信息不足以证明本机GPU算力，"
                "因此GPU算力可能性较低。"
            )
        return (
            f"探测响应显示{tool}接口或相关服务可访问，但未获得明确模型列表、"
            "本机模型运行状态或GPU证据，因此GPU算力可能性较低。"
        )

    return "未见直接GPU显存、GPU运行状态或本机模型加载证据，因此GPU算力可能性未知。"

def normalize_service_type(value: str) -> str:
    return clean(value) or "LLM服务"


def transform(df: pd.DataFrame, reported_single_ips: set[str]) -> pd.DataFrame:
    rows = []
    for _, row in df.iterrows():
        ip = clean(row.get("ip", ""))
        rows.append({
            "IP地址": ip,
            "端口号": row.get("port", ""),
            "协议类型": clean(row.get("protocol", "")),
            "服务类型": normalize_service_type(row.get("service_type", "")),
            "model_domain": clean(row.get("model_domain", "")),
            "模型部署工具": clean(row.get("deploy_tool", "")) or "未知",
            "工具版本": clean(row.get("deploy_version", "")),
            "模型信息": short_models(row.get("model_info", ""), row.get("evidence", "")),
            "访问链接": clean(row.get("link", "")),
            "探测说明": clean(row.get("evidence", "")),
            "证明截图": "",
            "存在GPU算力概率": gpu_label(row.get("gpu_likelihood", "未知")),
            "存在GPU算力的说明": gpu_note(row),
            "是否为上报IP": "是" if ip in reported_single_ips else "否",
            "_gpu_likelihood": clean(row.get("gpu_likelihood", "")),
            "_service_type": clean(row.get("service_type", "")),
        })
    return pd.DataFrame(rows)


def apply_reference_style(writer, sheet_name: str, row_count: int) -> None:
    ws = writer.sheets[sheet_name]
    thin = Side(style="thin", color="FF000000")
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    red_header = Font(name="宋体", size=11, bold=True, color="FFFF0000")
    url_header = Font(name="宋体", size=9, bold=True, color="FFFF0000")
    body_font = Font(name="宋体", size=11, bold=False, color="FF000000")
    url_font = Font(name="宋体", size=9, bold=False, color="FF000000")
    white_header = Font(name="宋体", size=11, bold=True, color="FFFFFFFF")
    blue_fill = PatternFill(fill_type="solid", fgColor="FF1F4E78")

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(OUT_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = url_header if col_idx == 9 else red_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if col_idx <= 13:
            cell.border = thin_border
        if col_idx == 14:
            cell.font = white_header
            cell.fill = blue_fill
            cell.alignment = Alignment(wrap_text=True)

    for row_idx in range(2, row_count + 2):
        ws.row_dimensions[row_idx].height = 130.5 if sheet_name == "GPU可能性高" else 110
        for col_idx in range(1, len(OUT_COLUMNS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = url_font if col_idx == 9 else body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col_idx <= 11:
                cell.border = thin_border


def write_result(df: pd.DataFrame, output: Path, reported_single_ips: set[str]) -> dict:
    normalized = transform(df, reported_single_ips)
    if normalized.empty:
        normalized = pd.DataFrame(columns=OUT_COLUMNS + ["_gpu_likelihood", "_service_type"])
    bucket_items = [
        ("GPU可能性高", normalized[normalized["_gpu_likelihood"].isin(["高", "high"])].copy()),
        ("GPU可能性中", normalized[normalized["_gpu_likelihood"].isin(["中", "medium"])].copy()),
        ("GPU可能性低", normalized[normalized["_gpu_likelihood"].isin(["低", "low"])].copy()),
    ]
    buckets = [(name, data) for name, data in bucket_items if not data.empty]
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        if not buckets:
            empty = pd.DataFrame(columns=OUT_COLUMNS)
            empty.to_excel(writer, index=False, sheet_name="无GPU可能性结果")
            apply_reference_style(writer, "无GPU可能性结果", 0)
        for sheet_name, sheet_df in buckets:
            sheet_df[OUT_COLUMNS].to_excel(writer, index=False, sheet_name=sheet_name)
            apply_reference_style(writer, sheet_name, len(sheet_df))
    counts = {name: len(data) for name, data in buckets}
    return {"gpu_high": counts.get("GPU可能性高", 0), "gpu_medium": counts.get("GPU可能性中", 0), "gpu_low": counts.get("GPU可能性低", 0), "output": str(output)}


def fill_screenshots(output: Path) -> None:
    script = Path(__file__).resolve().parent / "screenshot_evidence.py"
    if not script.is_file():
        print(f"Skip screenshots: missing {script}")
        return
    assets = output.parent / f"{output.stem}_screenshot_assets"
    command = [
        sys.executable, str(script), "fill",
        "--input", str(output), "--output", str(output), "--assets", str(assets),
    ]
    try:
        result = subprocess.run(command, check=False)
    except Exception as exc:
        print(f"Screenshot fill failed ({exc}); workbook kept at {output}")
        return
    if result.returncode != 0:
        print(f"Screenshot fill exited {result.returncode}; workbook kept at {output}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert raw LLM scan results into final GPU-likelihood workbook.")
    parser.add_argument("input", type=Path)
    parser.add_argument("-o", "--output", type=Path)
    parser.add_argument("--reported-source", type=Path, help="expanded *_result.csv used to mark single reported IPs")
    parser.add_argument("--no-screenshot", action="store_true", help="skip Chrome evidence screenshots after writing xlsx")
    args = parser.parse_args()
    src = args.input
    out = args.output or src.with_name(src.stem + "_最终版.xlsx")
    df = read_scan(src)
    stats = write_result(df, out, read_reported_single_ips(args.reported_source))
    print(f"Input rows: {len(df)}")
    print(f"GPU high: {stats['gpu_high']}")
    print(f"GPU medium: {stats['gpu_medium']}")
    print(f"GPU low: {stats['gpu_low']}")
    print(f"Output: {stats['output']}")
    if not args.no_screenshot:
        print("Filling evidence screenshots...")
        fill_screenshots(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
